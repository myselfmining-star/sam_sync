import io
import os
import pathlib
from typing import List, Union

import docx
import streamlit as st

# Add custom CSS to hide the GitHub icon
hide_github_icon = """
<style>
#GithubIcon {
  visibility: hidden;
}
</style>
"""
st.markdown(hide_github_icon, unsafe_allow_html=True)

# ==============================================================================
# SHARED CONSTANTS & HELPERS
# ==============================================================================
BASE_DIR = pathlib.Path(__file__).resolve().parent


def load_api_key() -> str:
    """Load Gemini API Key from 'API KEY.txt' if present, falling back to GEMINI_API_KEY."""
    key_file = BASE_DIR / "API KEY.txt"
    if key_file.exists():
        key = key_file.read_text(encoding="utf-8").strip()
        if key:
            return key
    return os.environ.get("GEMINI_API_KEY", "").strip()


def extract_raw_docx_text(source: Union[bytes, str, pathlib.Path]) -> str:
    """
    Read all raw paragraph and table text from an uploaded .docx file (bytes or path).
    Shared across Phase 1 and Phase 3.
    """
    doc_stream = io.BytesIO(source) if isinstance(source, bytes) else str(source)
    doc = docx.Document(doc_stream)
    chunks: List[str] = []

    for p in doc.paragraphs:
        txt = p.text.replace("\xa0", " ").strip()
        if txt:
            chunks.append(txt)

    for table in doc.tables:
        for row in table.rows:
            cells = [
                c.text.replace("\xa0", " ").strip()
                for c in row.cells
                if c.text.strip()
            ]
            if cells:
                chunks.append(" | ".join(cells))

    return "\n".join(chunks)


# Alias for backward compatibility with Phase 1
extract_docx_text = extract_raw_docx_text


# ==============================================================================
# UNIFIED STREAMLIT APPLICATION (PHASE 1 & PHASE 3 TABS)
# ==============================================================================
def run_combined_app() -> None:
    from main import render_phase1_tab
    from app import render_phase3_tab

    st.set_page_config(
        page_title="SAM VA Claim Engine — Phase 1 & Phase 3",
        page_icon="🎖️",
        layout="wide",
        initial_sidebar_state="expanded",
    )

    # Hide GitHub icon and apply shared styling
    st.markdown(
        """
        <style>
        #GithubIcon {
            visibility: hidden;
        }
        .viewerBadge_container__1QSob, .styles_viewerBadge__1yB5_, .viewerBadge_link__1S137, .viewerBadge_text__1JaDK {
            display: none !important;
        }
        .main-header {
            font-size: 2.1rem;
            font-weight: 700;
            color: #1e3a8a;
            margin-bottom: 0.2rem;
        }
        .sub-header {
            font-size: 1.05rem;
            color: #4b5563;
            margin-bottom: 1.0rem;
        }
        .stButton>button {
            border-radius: 6px;
            font-weight: 600;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="main-header">🎖️ SAM VA Claim Engine & Document Automation</div>',
        unsafe_allow_html=True,
    )
    st.markdown(
        '<div class="sub-header">Unified Workspace for Phase 1 (Intake & Claim Analysis) and Phase 3 (TP Health Document Package Generation)</div>',
        unsafe_allow_html=True,
    )

    # --------------------------------------------------------------------------
    # Shared Sidebar: Gemini AI Configuration
    # --------------------------------------------------------------------------
    st.sidebar.header("🔑 Gemini AI Configuration")

    detected_key = load_api_key()
    if detected_key:
        st.sidebar.success("✅ Gemini API Key detected automatically")

    gemini_api_key = st.sidebar.text_input(
        "Gemini API Key",
        value=detected_key,
        type="password",
        placeholder="Enter your Google Gemini API key...",
        help="Shared across Phase 1 and Phase 3 for Google GenAI extraction.",
    ).strip()

    model_options = [
        "gemini-3.5-flash",
        "gemini-3.5-flash-lite",
        "gemini-2.5-flash",
        "gemini-2.5-pro",
        "gemini-1.5-flash",
        "Custom",
    ]
    selected_model_option = st.sidebar.selectbox(
        "Gemini Model",
        options=model_options,
        index=0,
        help="Select the Gemini model used for document extraction.",
    )

    if selected_model_option == "Custom":
        gemini_model = (
            st.sidebar.text_input(
                "Enter Custom Model Name",
                value="gemini-3.5-flash",
            ).strip()
            or "gemini-3.5-flash"
        )
    else:
        gemini_model = selected_model_option

    # --------------------------------------------------------------------------
    # Phase Tabs
    # --------------------------------------------------------------------------
    tab_phase1, tab_phase3 = st.tabs(
        [
            "🎖️ Phase 1: Intake & Claim Analysis",
            "📋 Phase 3: TP Health Document Automation",
        ]
    )

    with tab_phase1:
        render_phase1_tab(api_key=gemini_api_key, model_name=gemini_model)

    with tab_phase3:
        render_phase3_tab(api_key=gemini_api_key, model_name=gemini_model)


if __name__ == "__main__":
    run_combined_app()

