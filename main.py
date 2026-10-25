import streamlit as st
import streamlit.runtime

hide_toolbar_buttons = """
<style>
/* Hide the entire toolbar actions container */
div[data-testid="stToolbarActions"] {
    display: none !important;
}

/* Fallback: Hide the specific header source action button */
button[data-testid="stBaseButton-header"] {
    display: none !important;
}
</style>
"""
st.markdown(hide_toolbar_buttons, unsafe_allow_html=True)

import os
import re
import sys
import io
import json
import zipfile
import tempfile
import argparse
import requests
from enum import Enum
from typing import Optional

from pydantic import BaseModel, Field
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_COLOR_INDEX
from google import genai
from google.genai import types

import pandas as pd

# -----------------------------------------------------------------------------
# 1. Database & Schema Initialization
# -----------------------------------------------------------------------------
API_BASE_URL = "https://script.google.com/macros/s/AKfycbwXKEBMtN71g_NRy9lO8ofEeH4NARaENPmvgRBG7351rRTZmMgyz_LBSDq5pPBTWr8S/exec"

def fetch_database_and_build_schema():
    print("Fetching live database and building schema...")
    try:
        def fetch_json(endpoint: str) -> dict:
            resp = requests.get(f"{API_BASE_URL}?data={endpoint}", timeout=15)
            resp.raise_for_status()
            return resp.json()

        question_bank_data = fetch_json("questions")
        condition_map = fetch_json("map")
        splits_data = fetch_json("splits")
    except Exception as e:
        raise RuntimeError(f"Failed to fetch live Google Sheets data: {e}")

    class ServiceConnectedCondition(BaseModel):
        raw_name: str = Field(description="Exact verbatim condition text from the source.")
        diagnostic_code: str = Field(description="The 4-digit VA diagnostic code as a string, e.g. '5243', '8620', '6260'.")
        rating: int = Field(description="Integer rating percentage (0, 10, 20, etc.).")
        date: Optional[str] = Field(default="", description="The effective date if present.")

    class DeniedCondition(BaseModel):
        raw_name: str = Field(description="Exact verbatim condition text from the source.")
        diagnostic_code: str = Field(description="The 4-digit VA diagnostic code as a string, e.g. '5243', '8620', '6260'.")
        date: Optional[str] = Field(default="", description="The effective date if present.")

    class VeteranIntakeData(BaseModel):
        full_name: str = Field(default="John Doe", description="The full name of the veteran.")
        deployments: list[str] = Field(description="List of combat/deployment locations.")
        current_sc_conditions: list[ServiceConnectedCondition]
        denied_conditions: list[DeniedCondition]
        
    print("Successfully loaded live database and schema.")
    return VeteranIntakeData, condition_map, splits_data, question_bank_data

def get_cached_database_and_schema():
    """Cached wrapper around fetch_database_and_build_schema for Streamlit."""
    return fetch_database_and_build_schema()

if streamlit.runtime.exists():
    get_cached_database_and_schema = st.cache_resource(ttl=3600, show_spinner=False)(get_cached_database_and_schema)


# -----------------------------------------------------------------------------
# 2. Helpers (Imported from shared module)
# -----------------------------------------------------------------------------
from streamlit_app import extract_docx_text, load_api_key



# -----------------------------------------------------------------------------
# 3. Extraction Routine (Gemini)
# -----------------------------------------------------------------------------
def extract_intake(file_path: str, model_name: str, dynamic_schema, api_key: Optional[str] = None):    
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found at '{file_path}'")

    effective_api_key = api_key or load_api_key()
    if not effective_api_key:
        raise ValueError("No API key found. Place it in 'API KEY.txt', set GEMINI_API_KEY, or specify it in settings.")

    client = genai.Client(api_key=effective_api_key)

    extraction_prompt = """
You are an expert VA Claims Intake and Rating Decision Parser.
Extract the veteran's full name, all service-connected conditions, previously denied/NSC conditions, and any deployment/combat locations.

Rules:
1. Extract the veteran's full name from the document (e.g., "Full Name: Brown"). If no name is found, default to "John Doe".
2. Extract the 4-digit VA diagnostic code from the brackets in each condition line into 'diagnostic_code' (e.g., [5243] -> "5243", [8620, left] -> "8620").
3. If a condition has a hyphenated diagnostic code in brackets (e.g., [5299 - 5297] or [5242 - 5243]), always extract the second 4-digit code ("5297" or "5243").
4. Strip the bracketed diagnostic code tag (e.g., [5243], [8620, left]), claim numbers/list numbers, ratings, and administrative dates (e.g., '11 Feb 2009') from 'raw_name' so 'raw_name' contains only the clean condition name text. Extract administrative dates strictly into the 'date' field.
5. If a condition is marked 'NSC', 'Non-Service Connected', or 'Denied', place it in denied_conditions. Never put it in current_sc_conditions.
6. For current_sc_conditions, extract the rating strictly as an integer (e.g., 20% becomes 20; 0% becomes 0).
7. Extract deployment locations (e.g., Vietnam, Iraq, Southwest Asia, etc.) into deployments.
"""

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".docx":
        print(f"Reading Word document: {file_path}")
        text_content = extract_docx_text(file_path)
        contents = [
            extraction_prompt,
            f"Document Contents:\n{text_content}"
        ]
    elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
        print(f"Reading image: {file_path}")
        mime_type = "image/jpeg" if ext in [".jpg", ".jpeg"] else f"image/{ext[1:]}"
        with open(file_path, "rb") as f:
            image_bytes = f.read()

        image_part = types.Part.from_bytes(data=image_bytes, mime_type=mime_type)
        contents = [
            extraction_prompt,
            image_part
        ]
    else:
        raise ValueError(f"Unsupported format: {ext}. Provide a .docx, .png, .jpg, or .jpeg file.")

    print(f"Querying Gemini ({model_name})...")
    response = client.models.generate_content(
        model=model_name,
        contents=contents,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=dynamic_schema,
            temperature=0.0
        )
    )

    parsed_json = json.loads(response.text)
    print("\n--- Extracted Veteran Data ---")
    print(json.dumps(parsed_json, indent=2))
    return parsed_json


def to_sentence_case(text: str) -> str:
    if not text:
        return ""
    s = text.strip()
    if not s:
        return ""
    # Capitalize the first character while keeping the rest intact
    return s[0].upper() + s[1:]

def clean_condition_name(raw_name: str) -> str:
    if not raw_name:
        return ""
    cleaned = re.sub(
        r"\s*\((?:also\s+)?claimed as [^)]*\)|,?\s*(?:also\s+)?claimed as .*$",
        "",
        str(raw_name),
        flags=re.IGNORECASE
    ).strip()
    return to_sentence_case(cleaned)

def normalize_dc(dc: str, raw_name: str = "") -> str:
    if not dc:
        return ""
    matches = re.findall(r"\d{4}", str(dc))
    if not matches:
        return str(dc).strip()
    code = int(matches[-1])
    name_lower = (raw_name or "").lower()

    # Spine routing (5235 <= code <= 5243) when raw_name is provided for routing/equivalence
    if name_lower and 5235 <= code <= 5243:
        if "cervical" in name_lower or "neck" in name_lower:
            return "5243"
        if any(k in name_lower for k in ("lumb", "thoracic", "back", "sacroiliac")):
            return "5237"

    # Peripheral nerve normalization
    if 8610 <= code <= 8630:
        code -= 100
    elif 8710 <= code <= 8730:
        code -= 200

    if code in (8510, 8511, 8512, 8513):
        return "8510"
    if code in (8520, 8521, 8524, 8526):
        return "8520"

    # Mental health general mapping: keep "9400" as its own code
    return str(code)

def format_percentage_range(val) -> str:
    if val is None or val == "":
        return "X%"
    if isinstance(val, (int, float)):
        if isinstance(val, float) and 0 < val < 1.0:
            return f"{round(val * 100)}%"
        return f"{int(val)}%"
    s = str(val).strip()
    try:
        f = float(s)
        if 0 < f < 1.0:
            return f"{round(f * 100)}%"
        elif f.is_integer() and not s.endswith("%"):
            return f"{int(f)}%"
    except ValueError:
        pass
    return s

def _combine_values(values: list[float]) -> float:
    combined = 0.0
    for val in sorted([v for v in values if v > 0], reverse=True):
        combined += (100.0 - combined) * (val / 100.0)
    return combined

def calculate_va_combined_rating(sc_conditions_or_ratings: list) -> tuple[int, int]:
    if not sc_conditions_or_ratings:
        return 0, 0

    # If passed plain integers (e.g., for the 50% simulation loop)
    if isinstance(sc_conditions_or_ratings[0], (int, float)):
        exact = round(_combine_values([float(r) for r in sc_conditions_or_ratings if r > 0]))
        return int(round(exact / 10.0) * 10), int(exact)

    # Otherwise, inspect dicts for paired Left/Right upper or lower extremities
    upper_left, upper_right, lower_left, lower_right, other_ratings = [], [], [], [], []
    upper_kw = ("upper", "shoulder", "elbow", "wrist", "arm", "hand", "cervical_radiculopathy", "radiculopathy_upper")
    lower_kw = ("lower", "hip", "knee", "ankle", "foot", "pes_planus", "plantar", "sciatic", "radiculopathy_lower")
    upper_dcs = ("5201", "5206", "5215", "8510", "8515")
    lower_dcs = ("5252", "5253", "5260", "5262", "5271", "5276", "5284", "8520")

    for sc in sc_conditions_or_ratings:
        r = sc.get("rating", 0)
        if r <= 0:
            continue
        norm_dc = normalize_dc(sc.get("diagnostic_code", ""), sc.get("raw_name", ""))
        text = f"{sc.get('raw_name', '')} {norm_dc}".lower()
        is_left = "left" in text and "right" not in text
        is_right = "right" in text and "left" not in text

        if (norm_dc in upper_dcs or any(k in text for k in upper_kw)) and (is_left or is_right):
            (upper_left if is_left else upper_right).append(float(r))
        elif (norm_dc in lower_dcs or any(k in text for k in lower_kw)) and (is_left or is_right):
            (lower_left if is_left else lower_right).append(float(r))
        else:
            other_ratings.append(float(r))

    bilateral_pool = []
    for left_list, right_list in [(upper_left, upper_right), (lower_left, lower_right)]:
        if left_list and right_list:
            bilateral_pool.extend(left_list + right_list)
        else:
            other_ratings.extend(left_list + right_list)

    if bilateral_pool:
        bi_combined = _combine_values(bilateral_pool)
        bi_with_factor = round(bi_combined * 1.10)  # Add 10% bilateral factor
        other_ratings.append(float(bi_with_factor))

    final_combined = round(_combine_values(other_ratings))
    rounded_10 = int(round(final_combined / 10.0) * 10)
    return rounded_10, int(final_combined)


# -----------------------------------------------------------------------------
# 4. Core Business Logic: Claim Analysis & Matching
# -----------------------------------------------------------------------------
def process_claim_data(parsed_json: dict, condition_map: dict, splits_data: dict):
    # 3. Calculate VA Math (with 10% Bilateral Factor)
    current_combined_rating, exact_cumulative = calculate_va_combined_rating(
        parsed_json.get("current_sc_conditions", [])
    )
    print(f"\nCalculated Current VA Rating: {current_combined_rating}% (Exact: {exact_cumulative}%)")

    # 4. Setup Data for Logic
    primary_mappings = condition_map.get("primary_mappings", {})
    increase_ranges = condition_map.get("typical_increase_ranges_master_table", {})
    
    def get_range(dc, raw_name=""):
        raw_val = increase_ranges.get(str(dc))
        if raw_val is None:
            raw_val = increase_ranges.get(normalize_dc(dc))
        if raw_val is None:
            raw_val = increase_ranges.get(normalize_dc(dc, raw_name), "X%")
        return format_percentage_range(raw_val)

    splits_lookup = {}
    for cat, items in splits_data.get("categories", {}).items():
        for item in items:
            combined_name = item.get("combined_name", "").lower().strip()
            splits_lookup[combined_name] = item

    potential_claims = []

    # 3. Add INCREASE (Intercept and Split if necessary)
    for sc in parsed_json.get("current_sc_conditions", []):
        if sc["rating"] < 100:  
            clean_sc_name = clean_condition_name(sc["raw_name"])
            raw_name_clean = clean_sc_name.lower()
            
            matched_split = splits_lookup.get(raw_name_clean) or next(
                (item for k, item in splits_lookup.items() if k and k in raw_name_clean), None
            )
            if matched_split:
                for target_key in ("split_target_a", "split_target_b"):
                    target = matched_split.get(target_key, {})
                    t_dc = str(target.get("diagnostic_code", "")).strip()
                    t_name = clean_condition_name(target.get("name", ""))
                    if t_name:
                        potential_claims.append({
                            "condition_name": t_name,
                            "diagnostic_code": normalize_dc(t_dc),
                            "claim_type": target.get("claim_type", "INCREASE"),
                            "current_rating": sc["rating"],
                            "linked_to": None,
                            "tier": 2,
                            "rating_range": format_percentage_range(target.get("rating_range")) if target.get("rating_range") else get_range(t_dc, t_name),
                        })
            else:
                # Normal condition (preserve actual extracted 4-digit DC without spine routing)
                potential_claims.append({
                    "condition_name": clean_sc_name,
                    "diagnostic_code": normalize_dc(sc["diagnostic_code"]),
                    "claim_type": "INCREASE",
                    "current_rating": sc["rating"],
                    "linked_to": None,
                    "tier": 2,
                    "rating_range": get_range(sc["diagnostic_code"], clean_sc_name)
                })
                
    # 4. Direct Secondaries
    JOINT_DCS = {"5003", "5201", "5206", "5215", "5237", "5242", "5243", "5252", "5253", "5257", "5260", "5262", "5271", "5276", "5284"}
    JOINT_KEYWORD_RE = re.compile(r"\b(knee|ankle|shoulder|back|neck|cervical|hip|elbow|wrist|lumbar)\b", re.IGNORECASE)
    for sc in parsed_json.get("current_sc_conditions", []):
        sc_dc = str(sc.get("diagnostic_code", "")).strip()
        clean_sc_name = clean_condition_name(sc.get("raw_name", ""))
        norm_sc_dc = normalize_dc(sc_dc, clean_sc_name)
        raw_lower = clean_sc_name.lower()
        
        is_musculoskeletal_or_missing = (not norm_sc_dc.isdigit()) or (5000 <= int(norm_sc_dc) <= 5399)
        is_joint = (norm_sc_dc in JOINT_DCS) or (is_musculoskeletal_or_missing and bool(JOINT_KEYWORD_RE.search(raw_lower)))

        if is_joint:
            joint_general_rule = primary_mappings.get("5003")
            if joint_general_rule:
                for sec in joint_general_rule.get("secondaries", []):
                    sec_dc = str(sec.get("diagnostic_code", "")).strip()
                    if sec_dc == "9434":
                        continue
                    sec_name = sec.get("condition_name", "Unknown")

                    db_nexus = sec.get("secondary_nexus_template", "Related to s/c [condition]")
                    final_nexus = (
                        db_nexus.replace("[condition]", raw_lower)
                        .replace("[specific joint] condition", raw_lower)
                    )
                    potential_claims.append({
                        "condition_name": sec_name,
                        "diagnostic_code": sec_dc,
                        "claim_type": sec.get("claim_type", "SECONDARY"),
                        "linked_to": final_nexus,
                        "primary_names": [clean_sc_name],
                        "tier": sec.get("priority_tier", 7),
                        "rating_range": format_percentage_range(sec.get("rating_range")) if sec.get("rating_range") else get_range(sec_dc, sec_name)
                    })

        rule = primary_mappings.get(norm_sc_dc) or primary_mappings.get(sc_dc)
        if not rule and norm_sc_dc.isdigit() and 9200 <= int(norm_sc_dc) <= 9599:
            rule = primary_mappings.get("9434")

        if rule:
            for sec in rule.get("secondaries", []):
                sec_name = sec.get("condition_name", "Unknown")
                sec_dc = str(sec.get("diagnostic_code", "")).strip()
                sec_dc_norm = normalize_dc(sec_dc, sec_name)
                
                # 1. Flip [OPPOSITE_SIDE] tag dynamically based on triggering primary vs active SC conditions
                if "[OPPOSITE_SIDE]" in sec_name:
                    primary_name_lower = raw_lower
                    
                    if "left" in primary_name_lower and "right" not in primary_name_lower:
                        opp_side = "Right"
                    elif "right" in primary_name_lower and "left" not in primary_name_lower:
                        opp_side = "Left"
                    else:
                        # Triggering primary condition is axial/neither (e.g., cervical or lumbar spine)
                        matching_sc_names = [
                            c["raw_name"].lower() for c in parsed_json.get("current_sc_conditions", [])
                            if normalize_dc(c.get("diagnostic_code", ""), c.get("raw_name", "")) == sec_dc_norm
                        ]
                        has_sc_left = any("left" in m for m in matching_sc_names)
                        has_sc_right = any("right" in m for m in matching_sc_names)
                        
                        if has_sc_left and not has_sc_right:
                            opp_side = "Right"
                        elif has_sc_right and not has_sc_left:
                            opp_side = "Left"
                        else:
                            opp_side = "Bilateral"
                    
                    sec_name = sec_name.replace("[OPPOSITE_SIDE]", opp_side)

                # 2. Pull exact nexus template from DB with dynamic nexus replacements
                db_nexus = sec.get("secondary_nexus_template", "Related to s/c [condition]")
                final_nexus = (
                    db_nexus.replace("[condition]", raw_lower)
                    .replace("[specific joint] condition", raw_lower)
                )

                potential_claims.append({
                    "condition_name": sec_name, 
                    "diagnostic_code": normalize_dc(sec_dc), 
                    "claim_type": sec.get("claim_type", "SECONDARY"),
                    "linked_to": final_nexus,
                    "primary_names": [clean_sc_name],
                    "tier": sec.get("priority_tier", 7),
                    "rating_range": format_percentage_range(sec.get("rating_range")) if sec.get("rating_range") else get_range(sec_dc, sec_name)
                })

    # 4.5 Category Rules: Chronic Physical Conditions to Mental Health ("5003" -> "9434")
    sc_names_list = []
    for c in parsed_json.get("current_sc_conditions", []):
        c_dc = normalize_dc(c.get("diagnostic_code", ""))
        c_rating = c.get("rating", 0)
        if c_dc.isdigit() and 6000 <= int(c_dc) <= 6099:
            continue
        if c_dc == "6100" and c_rating == 0:
            continue
        sc_names_list.append(clean_condition_name(c["raw_name"]))

    if sc_names_list:
        linked_str = "service-connected conditions"
        detailed_str_nexus = f"Related to chronic pain from service-connected {', '.join(sc_names_list)} which cause severe sleep disturbances"
        
        for sec in primary_mappings.get("5003", {}).get("secondaries", []):
            if str(sec.get("diagnostic_code", "")).strip() == "9434":
                db_nexus = sec.get("secondary_nexus_template", "Related to chronic pain from [condition]")
                final_nexus = db_nexus.replace("[condition]", linked_str)
                
                potential_claims.append({
                    "condition_name": sec.get("condition_name", "Unknown"),
                    "diagnostic_code": "9434",
                    "claim_type": sec.get("claim_type", "SECONDARY"),
                    "linked_to": final_nexus, # Concise for Section 3
                    "str_linked_to": detailed_str_nexus, # Detailed for Section 4
                    "primary_names": list(sc_names_list),
                    "tier": sec.get("priority_tier", 1),
                    "rating_range": format_percentage_range(sec.get("rating_range")) if sec.get("rating_range") else get_range("9434")
                })

    # 6. Dynamic Presumptives (100% DB-Driven via ExposurePresumptives Sheet)
    deployments = parsed_json.get("deployments", [])
    deploy_str = " ".join([d.lower() for d in deployments])
    presumptives_table = condition_map.get("exposure_presumptives", [])
    
    # Group matched rows and keywords by p_dc to aggregate multi-locations
    presumptive_matches = {}
    for row in presumptives_table:
        keyword = row.get("deployment_keyword", "").lower()
        if keyword and keyword in deploy_str:
            p_dc = str(row.get("diagnostic_code") or row.get("diagnostic codes") or "").strip()
            if not p_dc:
                continue
            if p_dc not in presumptive_matches:
                presumptive_matches[p_dc] = {
                    "first_row": row,
                    "keywords": set()
                }
            presumptive_matches[p_dc]["keywords"].add(keyword)

    for p_dc, match_info in presumptive_matches.items():
        first_row = match_info["first_row"]
        matched_keywords = match_info["keywords"]

        # Collect unique deployments matching any keyword for this p_dc (preserving order and casing)
        matched_locs = []
        for d in deployments:
            d_lower = d.lower()
            if any(kw in d_lower for kw in matched_keywords):
                if d not in matched_locs:
                    matched_locs.append(d)
        loc_display = " & ".join(matched_locs) if matched_locs else "qualifying locations"

        is_denied = any(
            normalize_dc(denied.get("diagnostic_code", ""), denied.get("raw_name", "")) == normalize_dc(p_dc)
            for denied in parsed_json.get("denied_conditions", [])
        )
        c_type = "REOPEN/PRESUMPTIVE" if is_denied else first_row.get("claim_type", "NEW CLAIM")

        db_nexus = first_row.get("nexus_reason", "Presumptive condition based on service records.")
        final_nexus = db_nexus.replace("[DEPLOYMENT_LOCATIONS]", loc_display)
        cond_name = first_row.get("condition_name", "Unknown")

        potential_claims.append({
            "condition_name": cond_name,
            "diagnostic_code": normalize_dc(p_dc),
            "claim_type": c_type,
            "linked_to": final_nexus,
            "tier": int(first_row.get("priority_tier") or 4),
            "rating_range": format_percentage_range(first_row.get("rating_range")) if first_row.get("rating_range") else get_range(p_dc, cond_name)
        })

    # 6.5 Presumptive/Reopen Tinnitus Rule when Tinnitus is not Service-Connected (Database-Driven via "6100")
    sc_norm_dcs = [normalize_dc(c.get("diagnostic_code", ""), c.get("raw_name", "")) for c in parsed_json.get("current_sc_conditions", [])]
    pot_norm_dcs = [normalize_dc(c.get("diagnostic_code", ""), c.get("condition_name", "")) for c in potential_claims]

    if "6260" not in sc_norm_dcs and "6260" not in pot_norm_dcs:
        denied_conds = parsed_json.get("denied_conditions", [])
        denied_norm_dcs_list = [normalize_dc(d.get("diagnostic_code", ""), d.get("raw_name", "")) for d in denied_conds]
        
        has_tinnitus_denied = ("6260" in denied_norm_dcs_list) or any("tinnitus" in d.get("raw_name", "").lower() for d in denied_conds)
        has_hearing_denied = ("6100" in denied_norm_dcs_list) or any("hearing" in d.get("raw_name", "").lower() for d in denied_conds)

        if has_tinnitus_denied or has_hearing_denied:
            tinnitus_rule = next(
                (sec for sec in primary_mappings.get("6100", {}).get("secondaries", []) if normalize_dc(sec.get("diagnostic_code", "")) == "6260"),
                None
            )
            if tinnitus_rule:
                base_nexus = tinnitus_rule.get("secondary_nexus_template") or tinnitus_rule.get("nexus_template") or ""
                if has_tinnitus_denied:
                    c_type = "REOPEN"
                    linked_nexus = f"Presumptive, {base_nexus}" if base_nexus and not base_nexus.lower().startswith("presumptive") else base_nexus
                else:
                    c_type = tinnitus_rule.get("claim_type", "NEW CLAIM — Presumptive")
                    linked_nexus = base_nexus

                potential_claims.append({
                    "condition_name": tinnitus_rule.get("condition_name", "Tinnitus"),
                    "diagnostic_code": normalize_dc(tinnitus_rule.get("diagnostic_code", "6260")),
                    "claim_type": c_type,
                    "linked_to": linked_nexus,
                    "tier": tinnitus_rule.get("priority_tier", 3),
                    "rating_range": format_percentage_range(tinnitus_rule.get("rating_range")) if tinnitus_rule.get("rating_range") else get_range("6260")
                })

    # 6.6 Rescue Denied (NSC) Conditions via Proposed Claims (2-Hop Reopen Rule)
    for d in parsed_json.get("denied_conditions", []):
        d_dc = normalize_dc(d.get("diagnostic_code", ""), d.get("raw_name", ""))
        if not d_dc:
            continue
        if any(normalize_dc(c.get("diagnostic_code", ""), c.get("condition_name", "")) == d_dc for c in potential_claims):
            continue
        if d_dc.isdigit() and ((6000 <= int(d_dc) <= 6099) or (7800 <= int(d_dc) <= 7805) or d_dc == "6100"):
            continue

        matched_rescue = False
        for parent_claim in list(potential_claims):
            parent_norm_dc = normalize_dc(parent_claim.get("diagnostic_code", ""), parent_claim.get("condition_name", ""))
            parent_rule = primary_mappings.get(parent_norm_dc)
            if not parent_rule:
                continue
            for sec in parent_rule.get("secondaries", []):
                if normalize_dc(sec.get("diagnostic_code", ""), sec.get("condition_name", "")) == d_dc:
                    potential_claims.append({
                        "condition_name": sec.get("condition_name", clean_condition_name(d.get("raw_name", "Unknown"))),
                        "diagnostic_code": normalize_dc(sec.get("diagnostic_code", d_dc)),
                        "claim_type": "REOPEN/SECONDARY",
                        "linked_to": f"Related to {parent_claim['condition_name'].lower()} when it is service connected",
                        "primary_names": [parent_claim["condition_name"]],
                        "tier": sec.get("priority_tier", 3),
                        "rating_range": format_percentage_range(sec.get("rating_range")) if sec.get("rating_range") else get_range(d_dc),
                    })
                    matched_rescue = True
                    break
            if matched_rescue:
                break

    # 6.7 Chained Secondaries off Proposed Tinnitus ("6260") When Direct Physical Secondaries Are Empty
    has_direct_physical_secondary = any(
        c.get("claim_type") == "SECONDARY" and normalize_dc(c.get("diagnostic_code", "")) != "9434"
        for c in potential_claims
    )
    has_proposed_tinnitus = ("6260" not in sc_norm_dcs) and any(
        normalize_dc(c.get("diagnostic_code", ""), c.get("condition_name", "")) == "6260"
        for c in potential_claims
    )
    if not has_direct_physical_secondary and has_proposed_tinnitus:
        for sec in primary_mappings.get("6260", {}).get("secondaries", []):
            sec_dc = str(sec.get("diagnostic_code", "")).strip()
            potential_claims.append({
                "condition_name": sec.get("condition_name", "Unknown"),
                "diagnostic_code": normalize_dc(sec_dc),
                "claim_type": "SECONDARY",
                "linked_to": "Related to tinnitus when it is service connected",
                "primary_names": ["Tinnitus"],
                "tier": sec.get("priority_tier", 7),
                "rating_range": format_percentage_range(sec.get("rating_range")) if sec.get("rating_range") else get_range(sec_dc),
            })

    # 7. Check against Denied Conditions (The Reopen Filter)
    denied_names = [clean_condition_name(d["raw_name"]).lower() for d in parsed_json.get("denied_conditions", [])]
    denied_norm_dcs = {
        normalize_dc(d.get("diagnostic_code", ""), d.get("raw_name", ""))
        for d in parsed_json.get("denied_conditions", [])
        if d.get("diagnostic_code")
    }
    has_denied_mental = any(d_dc.isdigit() and 9200 <= int(d_dc) <= 9599 for d_dc in denied_norm_dcs)

    for claim in potential_claims:
        # Objective Rule: Active service-connected conditions are INCREASEs, not REOPENs
        if claim["claim_type"] == "INCREASE":
            continue
            
        norm_dc = normalize_dc(claim.get("diagnostic_code", ""), claim.get("condition_name", ""))
        is_mental_claim = norm_dc.isdigit() and 9200 <= int(norm_dc) <= 9599
        
        if claim["condition_name"].strip().lower() in denied_names or norm_dc in denied_norm_dcs or (is_mental_claim and has_denied_mental):
            if "PRESUMPTIVE" not in claim["claim_type"] and "REOPEN" not in claim["claim_type"]:
                claim["claim_type"] = "REOPEN/SECONDARY"

    # 1. Exclusions: Vision/Eye, Scars, One-Mental-Health Rule, and Maxed/Fixed/0%-Hearing INCREASEs
    has_sc_mental = any(d.isdigit() and 9200 <= int(d) <= 9599 for d in sc_norm_dcs)
    filtered_claims = []
    for c in potential_claims:
        norm_dc = normalize_dc(c.get("diagnostic_code", ""), c.get("condition_name", ""))
        c_name_lower = c.get("condition_name", "").lower()

        if norm_dc.isdigit() and 6000 <= int(norm_dc) <= 6099:
            continue
        if (norm_dc.isdigit() and 7800 <= int(norm_dc) <= 7805) or "scar" in c_name_lower:
            continue
        if c["claim_type"] != "INCREASE" and norm_dc.isdigit() and 9200 <= int(norm_dc) <= 9599 and has_sc_mental:
            continue
        if c["claim_type"] == "INCREASE" and (
            norm_dc == "6260"
            or (norm_dc == "6100" and c.get("current_rating", 0) == 0)
            or "tinnitus" in c_name_lower
            or "fixed" in str(c.get("rating_range", "")).lower()
        ):
            continue
        filtered_claims.append(c)

    # 2. Prevent NEW/SECONDARY claims for conditions the veteran already has (contralateral side check)
    valid_claims = []
    for claim in filtered_claims:
        if claim["claim_type"] != "INCREASE":
            norm_dc = normalize_dc(claim.get("diagnostic_code", ""), claim.get("condition_name", ""))
            matching_sc = [
                sc for sc in parsed_json.get("current_sc_conditions", [])
                if normalize_dc(sc.get("diagnostic_code", ""), sc.get("raw_name", "")) == norm_dc
            ]
            if matching_sc:
                claim_name_lower = claim.get("condition_name", "").lower()
                claim_has_left = "left" in claim_name_lower
                claim_has_right = "right" in claim_name_lower

                sc_names_lower = [sc.get("raw_name", "").lower() for sc in matching_sc]
                sc_has_left = any("left" in name for name in sc_names_lower)
                sc_has_right = any("right" in name for name in sc_names_lower)
                sc_has_bilateral = any("bilateral" in name for name in sc_names_lower)

                if claim_has_right and not claim_has_left and (sc_has_left and not sc_has_right and not sc_has_bilateral):
                    pass # Keep contralateral claim
                elif claim_has_left and not claim_has_right and (sc_has_right and not sc_has_left and not sc_has_bilateral):
                    pass # Keep contralateral claim
                else:
                    claim_words = set(re.findall(r"[a-z]{4,}", clean_condition_name(claim_name_lower)))
                    shares_name = any(
                        bool(claim_words & set(re.findall(r"[a-z]{4,}", clean_condition_name(sc_name))))
                        or clean_condition_name(sc_name) in clean_condition_name(claim_name_lower)
                        or clean_condition_name(claim_name_lower) in clean_condition_name(sc_name)
                        for sc_name in sc_names_lower
                    )
                    if norm_dc != "6204" or shares_name:
                        continue # Skip it, they already have it
        valid_claims.append(claim)

    # 3. Merge duplicate claims (Pyramiding) without corrupting Presumptive vs Secondary nexus
    merged_claims = {}
    for claim in valid_claims:
        dc = claim.get("diagnostic_code", "")
        ctype = claim.get("claim_type", "")
        
        if ctype == "INCREASE":
            merge_key = f"INCREASE_{claim['condition_name'].strip().lower()}"
        else:
            c_name_lower = claim.get("condition_name", "").lower()
            side_tag = "left" if "left" in c_name_lower else ("right" if "right" in c_name_lower else "any")
            merge_key = f"NON_INC_{normalize_dc(dc, claim.get('condition_name', ''))}_{side_tag}"

        if merge_key in merged_claims:
            existing = merged_claims[merge_key]
            existing_ctype = existing.get("claim_type", "")
            existing_tier = int(existing.get("tier") or 7)
            new_tier = int(claim.get("tier") or 7)

            both_secondary = ("SECONDARY" in existing_ctype) and ("SECONDARY" in ctype)
            is_new_presumptive = ("PRESUMPTIVE" in ctype) or (ctype == "NEW CLAIM")
            is_existing_presumptive = ("PRESUMPTIVE" in existing_ctype) or (existing_ctype == "NEW CLAIM")

            if both_secondary:
                existing_link = existing.get("linked_to", "")
                new_link = claim.get("linked_to", "")
                if new_link and new_link not in existing_link:
                    clean_new = (
                        new_link.replace("Related to NSAID use for s/c ", "")
                        .replace("NSAID use for s/c ", "")
                        .replace("Related to s/c ", "")
                        .replace("Related to ", "")
                    )
                    existing["linked_to"] = f"{existing_link} and {clean_new}"
                existing_primaries = list(existing.get("primary_names", []))
                for p_name in claim.get("primary_names", []):
                    if p_name and p_name.lower() not in [ep.lower() for ep in existing_primaries]:
                        existing_primaries.append(p_name)
                existing["primary_names"] = existing_primaries
                existing["tier"] = min(existing_tier, new_tier)
                if "REOPEN" in ctype and "REOPEN" not in existing_ctype:
                    existing["claim_type"] = ctype
                if claim.get("str_linked_to") and not existing.get("str_linked_to"):
                    existing["str_linked_to"] = claim["str_linked_to"]
            elif is_new_presumptive and not is_existing_presumptive:
                merged_claims[merge_key] = claim
            elif is_existing_presumptive and not is_new_presumptive:
                pass
            elif new_tier < existing_tier:
                merged_claims[merge_key] = claim
        else:
            merged_claims[merge_key] = claim

    final_claims = list(merged_claims.values())

    # 4. Sort the claims by their assigned priority tier (safely cast tier to int)
    final_claims.sort(key=lambda x: int(x.get("tier") or 99))

    return current_combined_rating, exact_cumulative, final_claims



# -----------------------------------------------------------------------------
# 5. Document Builder (Word Output)
# -----------------------------------------------------------------------------
def generate_docx_report(parsed_json: dict, current_combined_rating: int, exact_cumulative: int, final_claims: list[dict], condition_templates: dict):
    doc = Document()
    veteran_name = parsed_json.get("full_name", "John Doe")
    
    # Set default global font to Calibri 12pt
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(12)

    # --- SECTION 1: HEADER & MATH ---
    doc.add_paragraph(f"Good morning {veteran_name},")
    doc.add_paragraph("After reviewing your eBenefits snapshot, I have the following analysis of your VA Disabilities:")
    
    # 1. Use actual bullet points for the intro math
    doc.add_paragraph(f"You're currently {current_combined_rating}% (Cumulative {exact_cumulative}%)", style='List Bullet')
    if 60 <= current_combined_rating < 100:
        raw_ratings = [condition["rating"] for condition in parsed_json.get("current_sc_conditions", [])]
        count_50s = 0
        simulated_ratings = list(raw_ratings)
        while calculate_va_combined_rating(simulated_ratings)[0] < 100 and count_50s < 20:
            simulated_ratings.append(50)
            count_50s += 1
        doc.add_paragraph(f"You will require ~{count_50s} additional 50% conditions or similar combination to receive a VA Disability Rating of 100%", style='List Bullet')
    
    # 2. Rating Block
    for label, val in [
        ("Rating:", " Yes"),
        ("Percentage:", f" {current_combined_rating}% (cumulative ~{exact_cumulative}%)"),
        ("Type of Claim:", " Appeal/New")
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(label).bold = True
        p.add_run(val)

    # --- SECTION 2: CURRENT RATING ---
    p_header = doc.add_paragraph()
    run_header = p_header.add_run("CURRENT RATING")
    run_header.bold = True
    run_header.font.color.rgb = RGBColor(0, 0, 0)

    def add_numbered_condition(idx, name, status_text, date_val):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(0)
        date_str = f" ({date_val})" if date_val else ""
        formatted_name = to_sentence_case(name) if (name and (name.islower() or name.strip()[0].islower())) else name
        p.add_run(f"{idx}) ")
        p.add_run(formatted_name).bold = True
        p.add_run(f" — {status_text}{date_str}")

    curr_sc = parsed_json.get("current_sc_conditions", [])
    denied = parsed_json.get("denied_conditions", [])

    for i, sc in enumerate(curr_sc, 1):
        add_numbered_condition(i, sc["raw_name"], f"{sc['rating']}%", sc.get("date"))

    for i, d in enumerate(denied, len(curr_sc) + 1):
        add_numbered_condition(i, d["raw_name"], "NSC", d.get("date"))

    # --- SECTION 3: POTENTIAL CONDITIONS ---
    doc.add_paragraph()
    p_pot_head = doc.add_paragraph()
    run_pot_head = p_pot_head.add_run("POTENTIAL CONDITIONS")
    run_pot_head.bold = True
    run_pot_head.font.color.rgb = RGBColor(0, 0, 0)

    for claim in final_claims:
        c_type = claim['claim_type']
        
        if c_type == "INCREASE":
            c_type = f"INCREASE (currently {claim.get('current_rating', 0)}%)"
            linked_str = ""
        else:
            link = claim.get('linked_to', '')
            linked_str = f" — {link}" if link else ""

        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{claim['condition_name']} — {c_type}{linked_str} ")
        
        # Apply Blue Bold to the extracted percentage range
        pot_range = claim.get('rating_range', 'X%')
        pot_run = p.add_run(f"(Potentially {pot_range})")
        pot_run.bold = True
        pot_run.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)

    # --- SECTION 4: STR EVALUATION ---
    doc.add_paragraph()
    p_str_head = doc.add_paragraph()
    run_str_head = p_str_head.add_run("SERVICE TREATMENT RECORDS (STR) EVALUATION")
    run_str_head.bold = True
    run_str_head.font.color.rgb = RGBColor(0, 0, 0)
    
    # STRs Reviewed line (Red)
    p_str = doc.add_paragraph()
    run_str = p_str.add_run("STRs WERE REVIEWED")
    run_str.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    for idx, claim in enumerate(final_claims, 1):
        p = doc.add_paragraph()
        
        c_type = claim['claim_type']
        link = claim.get('str_linked_to') or claim.get('linked_to', '')
        
        # Number the conditions and separate bold/unbold text
        if c_type == "INCREASE":
            bold_text = f"\n{idx}) {claim['condition_name']} — {c_type}"
            unbold_text = ""
        else:
            bold_text = f"\n{idx}) {claim['condition_name']} — {c_type}"
            unbold_text = f" — {link}" if link else ""

        p.add_run(bold_text).bold = True
        if unbold_text:
            primary_names = [pn for pn in claim.get("primary_names", []) if pn]
            if primary_names:
                sorted_pns = sorted(primary_names, key=len, reverse=True)
                pattern = re.compile("(" + "|".join(re.escape(pn) for pn in sorted_pns) + ")", re.IGNORECASE)
                for part in pattern.split(unbold_text):
                    if not part:
                        continue
                    r = p.add_run(part)
                    if pattern.fullmatch(part):
                        r.bold = True
                        r.italic = True
            else:
                p.add_run(unbold_text)

        # Lookup template with fallback to normalized Diagnostic Code
        claim_dc = str(claim.get("diagnostic_code", "")).strip()
        template = condition_templates.get(claim_dc, {})
        if not template or not template.get("questions"):
            norm_dc = normalize_dc(claim_dc, claim.get("condition_name", ""))
            if norm_dc in condition_templates:
                fallback_tmpl = condition_templates[norm_dc]
                if fallback_tmpl.get("questions"):
                    template = fallback_tmpl
        
        if not template or (not template.get("questions") and not template.get("symptom_checklist")):
            p_miss = doc.add_paragraph(style='List Bullet')
            run_miss = p_miss.add_run(f"[MISSING] No mapping found in question bank for Diagnostic Code: '{claim_dc}'")
            run_miss.font.highlight_color = WD_COLOR_INDEX.YELLOW
            continue


        # Print Standard Questions (Yellow Highlight)
        for q in template.get("questions", []):
            if isinstance(q, str):
                p_q = doc.add_paragraph(style='List Bullet')
                run_q = p_q.add_run(q)
                run_q.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif isinstance(q, dict) and "question_text" in q:
                p_q = doc.add_paragraph(style='List Bullet')
                run_q = p_q.add_run(q["question_text"])
                run_q.font.highlight_color = WD_COLOR_INDEX.YELLOW
                for sub in q.get("sub_bullets", []):
                    p_sub = doc.add_paragraph(style='List Bullet 2')
                    run_sub = p_sub.add_run(sub)
                    run_sub.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    
        # Print Symptom Checklist
        if "symptom_checklist_header" in template:
            p_symp_head = doc.add_paragraph(style='List Bullet')
            run_symp_head = p_symp_head.add_run(template["symptom_checklist_header"])
            run_symp_head.font.highlight_color = WD_COLOR_INDEX.YELLOW
            
        if "symptom_checklist" in template:
            for symp in template["symptom_checklist"]:
                doc.add_paragraph(symp, style='List Bullet 2')
                
        # Print Closing Block (Colored per Rules)
        if "closing_block" in template:
            for close_text in template["closing_block"]:
                p_close = doc.add_paragraph(style='List Bullet')
                run_close = p_close.add_run(close_text)
                
                # Apply specific colors based on exception rules
                if "sleep study" in close_text.lower() and "require" in close_text.lower():
                    run_close.font.color.rgb = RGBColor(0xFF, 0x00, 0x00) # Red
                    run_close.bold = True
                elif "gerd" in close_text.lower() and "appropriate test" in close_text.lower():
                    run_close.font.color.rgb = RGBColor(0xED, 0x5C, 0x57) # Pink-Red
                    run_close.bold = True
                    run_close.italic = True
                else:
                    run_close.font.color.rgb = RGBColor(0x00, 0x00, 0xFF) # Blue

    # Save Document bytes
    safe_name = veteran_name.replace(" ", "_")
    output_filename = f"VA_Disability_Analysis_{safe_name}_Final.docx"
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    docx_bytes = buffer.getvalue()
    
    return doc, output_filename, docx_bytes


# -----------------------------------------------------------------------------
# 6. High-Level Document Processor
# -----------------------------------------------------------------------------
def process_single_file(file_path: str, filename: str, model_name: str, dynamic_schema, condition_map, splits_data, question_bank_data, api_key: Optional[str] = None) -> dict:
    parsed_json = extract_intake(file_path, model_name, dynamic_schema, api_key=api_key)
    current_combined_rating, exact_cumulative, final_claims = process_claim_data(parsed_json, condition_map, splits_data)
    condition_templates = question_bank_data.get("condition_templates", {})
    doc, output_filename, docx_bytes = generate_docx_report(parsed_json, current_combined_rating, exact_cumulative, final_claims, condition_templates)
    
    return {
        "input_name": filename,
        "veteran_name": parsed_json.get("full_name", "John Doe"),
        "output_filename": output_filename,
        "docx_bytes": docx_bytes,
        "doc_object": doc,
        "current_combined_rating": current_combined_rating,
        "exact_cumulative": exact_cumulative,
        "parsed_json": parsed_json,
        "final_claims": final_claims,
        "status": "success",
        "error": None
    }


# -----------------------------------------------------------------------------
# 7. Streamlit Tab Interface (Phase 1)
# -----------------------------------------------------------------------------
def render_phase1_tab(api_key: str, model_name: str) -> None:
    st.header("🎖️ Phase 1 — Intake Document Analysis & Claim Engine")
    st.caption(
        "Batch Intake Document Analysis, VA Combined Rating Calculation & Automated Report Generation"
    )

    st.sidebar.divider()
    st.sidebar.subheader("📊 Phase 1: Live Rule Database")

    if st.sidebar.button("🔄 Reload Database Rules", key="phase1_reload_db_btn"):
        st.cache_resource.clear()
        st.sidebar.success("Cache cleared! Live Google Sheets database will reload on next request.")

    # Initialize session state for batch results
    if "batch_results" not in st.session_state:
        st.session_state["batch_results"] = []

    # --- MAIN CONTENT AREA ---
    st.subheader("1. Upload Batch of Intake Documents")
    uploaded_files = st.file_uploader(
        "Choose .docx intake forms or screenshot code sheets (.png, .jpg, .jpeg, .webp)",
        type=["docx", "png", "jpg", "jpeg", "webp"],
        accept_multiple_files=True,
        key="phase1_intake_uploader",
        help="You can drag and drop multiple files at once.",
    )

    if uploaded_files:
        st.info(f"📂 **{len(uploaded_files)} file(s) selected for processing.**")

        col_btn1, col_btn2, _ = st.columns([2, 1.5, 4])
        start_processing = col_btn1.button(
            "🚀 Process Batch Documents",
            type="primary",
            use_container_width=True,
            key="phase1_process_btn",
        )
        clear_btn = col_btn2.button(
            "🗑️ Clear Results",
            use_container_width=True,
            key="phase1_clear_btn",
        )

        if clear_btn:
            st.session_state["batch_results"] = []
            st.rerun()

        if start_processing:
            if not api_key.strip():
                st.error("Please enter a valid Gemini API Key in the sidebar or add it to 'API KEY.txt'.")
                return

            # Fetch DB and schema
            with st.spinner("Connecting to live database and building condition schemas..."):
                try:
                    dynamic_schema, condition_map, splits_data, question_bank_data = get_cached_database_and_schema()
                except Exception as e:
                    st.error(f"Failed to fetch live database from Google Sheets: {e}")
                    return

            total_files = len(uploaded_files)
            progress_bar = st.progress(0)
            status_placeholder = st.empty()
            results = []

            for i, uploaded_file in enumerate(uploaded_files):
                progress = (i) / total_files
                progress_bar.progress(progress)
                status_placeholder.markdown(f"⏳ **Processing ({i + 1}/{total_files}):** `{uploaded_file.name}` ...")

                suffix = os.path.splitext(uploaded_file.name)[1]
                temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
                try:
                    temp_file.write(uploaded_file.getbuffer())
                    temp_file.close()
                    tmp_path = temp_file.name

                    # Run full processing pipeline
                    res = process_single_file(
                        file_path=tmp_path,
                        filename=uploaded_file.name,
                        model_name=model_name,
                        dynamic_schema=dynamic_schema,
                        condition_map=condition_map,
                        splits_data=splits_data,
                        question_bank_data=question_bank_data,
                        api_key=api_key.strip()
                    )
                    results.append(res)
                except Exception as e:
                    print(f"Error processing {uploaded_file.name}: {e}")
                    results.append({
                        "input_name": uploaded_file.name,
                        "veteran_name": "Error",
                        "output_filename": f"Error_{uploaded_file.name}.txt",
                        "docx_bytes": None,
                        "doc_object": None,
                        "current_combined_rating": 0,
                        "exact_cumulative": 0,
                        "parsed_json": {},
                        "final_claims": [],
                        "status": "error",
                        "error": str(e)
                    })
                finally:
                    if os.path.exists(tmp_path):
                        try:
                            os.remove(tmp_path)
                        except Exception:
                            pass

            progress_bar.progress(1.0)
            status_placeholder.empty()

            st.session_state["batch_results"] = results
            st.rerun()

    # --- DISPLAY RESULTS & ZIP DOWNLOAD ---
    batch_results = st.session_state.get("batch_results", [])
    if batch_results:
        st.divider()
        st.subheader("2. Analysis Results & Downloads")

        success_results = [r for r in batch_results if r["status"] == "success"]
        failed_results = [r for r in batch_results if r["status"] == "error"]

        # Top summary metrics
        m1, m2, m3 = st.columns(3)
        m1.metric("Total Files Processed", len(batch_results))
        m2.metric("Successful Analyses", len(success_results))
        m3.metric("Failed Analyses", len(failed_results))

        # ZIP Download Button
        if success_results:
            zip_buffer = io.BytesIO()
            used_filenames = set()

            with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as zf:
                for res in success_results:
                    base_name = res["output_filename"]
                    out_name = base_name
                    counter = 1
                    while out_name in used_filenames:
                        name_part, ext_part = os.path.splitext(base_name)
                        out_name = f"{name_part}_{counter}{ext_part}"
                        counter += 1
                    used_filenames.add(out_name)
                    zf.writestr(out_name, res["docx_bytes"])

            zip_buffer.seek(0)
            zip_data = zip_buffer.getvalue()

            st.download_button(
                label=f"📦 Download All {len(success_results)} Reports as ZIP",
                data=zip_data,
                file_name="VA_Disability_Analysis_Reports.zip",
                mime="application/zip",
                type="primary",
                use_container_width=True,
                key="phase1_zip_dl_btn",
            )

        # Show detailed result cards
        st.markdown("### Processed Veterans")
        for i, res in enumerate(batch_results):
            if res["status"] == "success":
                vet_name = res["veteran_name"]
                comb = res["current_combined_rating"]
                exact = res["exact_cumulative"]
                in_name = res["input_name"]
                claims_count = len(res["final_claims"])
                sc_count = len(res["parsed_json"].get("current_sc_conditions", []))
                denied_count = len(res["parsed_json"].get("denied_conditions", []))

                with st.expander(f"👤 {vet_name} — Current Rating: {comb}% (Exact: {exact}%) — 📄 {in_name}", expanded=(i == 0)):
                    col_dl, col_m1, col_m2, col_m3 = st.columns([2, 1.2, 1.2, 1.2])
                    
                    col_dl.download_button(
                        label=f"⬇️ Download {res['output_filename']}",
                        data=res["docx_bytes"],
                        file_name=res["output_filename"],
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key=f"dl_{i}_{res['output_filename']}"
                    )
                    col_m1.metric("Service Connected", sc_count)
                    col_m2.metric("Denied / NSC", denied_count)
                    col_m3.metric("Potential Claims", claims_count)

                    tab_claims, tab_sc, tab_denied, tab_json = st.tabs([
                        f"Potential Claims ({claims_count})",
                        f"Current SC ({sc_count})",
                        f"Denied ({denied_count})",
                        "Raw Extracted Data"
                    ])

                    with tab_claims:
                        if res["final_claims"]:
                            claims_rows = []
                            for idx, c in enumerate(res["final_claims"], 1):
                                claims_rows.append({
                                    "#": idx,
                                    "Condition": c.get("condition_name", "Unknown"),
                                    "Diagnostic Code": c.get("diagnostic_code", ""),
                                    "Claim Type": c.get("claim_type", ""),
                                    "Potential Rating": c.get("rating_range", "X%"),
                                    "Tier": c.get("tier", ""),
                                    "Nexus / Link": c.get("linked_to") or "Primary / Direct"
                                })
                            st.dataframe(pd.DataFrame(claims_rows), use_container_width=True, hide_index=True)
                        else:
                            st.write("No potential claims identified.")

                    with tab_sc:
                        sc_list = res["parsed_json"].get("current_sc_conditions", [])
                        if sc_list:
                            sc_rows = []
                            for sc in sc_list:
                                sc_rows.append({
                                    "Condition": sc.get("raw_name"),
                                    "Rating": f"{sc.get('rating')}%",
                                    "Date": sc.get("date") or "N/A",
                                    "Diagnostic Code": str(sc.get("diagnostic_code", ""))
                                })
                            st.dataframe(pd.DataFrame(sc_rows), use_container_width=True, hide_index=True)
                        else:
                            st.write("No service-connected conditions found.")

                    with tab_denied:
                        denied_list = res["parsed_json"].get("denied_conditions", [])
                        if denied_list:
                            den_rows = []
                            for d in denied_list:
                                den_rows.append({
                                    "Condition": d.get("raw_name"),
                                    "Status": "NSC / Denied",
                                    "Date": d.get("date") or "N/A",
                                    "Diagnostic Code": str(d.get("diagnostic_code", ""))
                                })
                            st.dataframe(pd.DataFrame(den_rows), use_container_width=True, hide_index=True)
                        else:
                            st.write("No denied conditions found.")

                    with tab_json:
                        st.json(res["parsed_json"])
            else:
                with st.expander(f"❌ Failed: {res['input_name']}", expanded=True):
                    st.error(f"Error processing `{res['input_name']}`:\n\n{res['error']}")


# -----------------------------------------------------------------------------
# 8. Command Line Interface (CLI fallback)
# -----------------------------------------------------------------------------
def run_cli():
    parser = argparse.ArgumentParser(description="Extract VA Intake data using Gemini.")
    parser.add_argument("file_path", nargs="?", default=None, help="Path to .docx intake form or code sheet screenshot (.png/.jpg)")
    parser.add_argument("--model", default="gemini-3.5-flash-lite", help="Gemini model name")
    args = parser.parse_args()

    if not args.file_path:
        print("\n[SAM VA Claim Engine]")
        print("To run the combined Streamlit web interface:")
        print("    streamlit run streamlit_app.py")
        print("\nTo run via command line:")
        print("    python main.py <file_path> [--model <model_name>]")
        sys.exit(0)

    # 1. Fetch DB and build schema FIRST
    DynamicVeteranSchema, condition_map, splits_data, question_bank_data = fetch_database_and_build_schema()
    
    # 2. Extract Data & Run Analysis
    res = process_single_file(
        file_path=args.file_path,
        filename=os.path.basename(args.file_path),
        model_name=args.model,
        dynamic_schema=DynamicVeteranSchema,
        condition_map=condition_map,
        splits_data=splits_data,
        question_bank_data=question_bank_data
    )

    # Save to disk for CLI
    res["doc_object"].save(res["output_filename"])
    print(f"\nPipeline Complete! Word Document saved as '{res['output_filename']}'")


# -----------------------------------------------------------------------------
# 9. Main Entrypoint
# -----------------------------------------------------------------------------
if __name__ == "__main__":
    if streamlit.runtime.exists():
        from streamlit_app import run_combined_app
        run_combined_app()
    else:
        run_cli()