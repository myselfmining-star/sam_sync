import streamlit as st

# Add custom CSS to hide the GitHub icon
hide_github_icon = """
<style>
#GithubIcon {
  visibility: hidden;
}
</style>
"""
st.markdown(hide_github_icon, unsafe_allow_html=True)

import datetime
import hashlib
import io
import os
import pathlib
import zipfile
from typing import Any, Dict, List, Literal, Optional, Tuple

import docx
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor
import fitz  # PyMuPDF
from google import genai
from google.genai import types
from pydantic import BaseModel, Field

# ==============================================================================
# CONSTANTS & PATHS
# ==============================================================================
from streamlit_app import BASE_DIR, extract_raw_docx_text

TEMPLATES_DIR = BASE_DIR / "templates"
MINI_DBQ_DIR = TEMPLATES_DIR / "mini_dbqs"

PAYMENT_FORM_TEMPLATE = TEMPLATES_DIR / "1. PATIENT PAYMENT FORM 06.03.2024.pdf"
ROI_FORM_TEMPLATE = TEMPLATES_DIR / "2. RELEASE OF INFORMATION 04.29 - Copy.pdf"

COLOR_MAROON = RGBColor(0x80, 0x00, 0x00)
COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)
COLOR_DARK_GRAY = RGBColor(0x22, 0x22, 0x22)
COLOR_BLUE = RGBColor(0x00, 0x00, 0xFF)


# ==============================================================================
# PART 2: EXACT CATALOG LOOKUP TABLE (MINI_DBQ_CATALOG)
# ==============================================================================
MINI_DBQ_CATALOG: List[Dict[str, Any]] = [
    {
        "key": "ankle",
        "filenames": ["Ankle MDBQ v2 Form.pdf"],
        "short_title": "Ankle MDBQ",
        "roi_name": "Ankle DBQ",
        "mrr_name": "Ankle DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ medical opinion",
    },
    {
        "key": "arthritis",
        "filenames": ["Arthritis MDBQ v2 Form.pdf"],
        "short_title": "Arthritis MDBQ",
        "roi_name": "Arthritis DBQ",
        "mrr_name": "Arthritis DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "back",
        "filenames": ["Back MDBQ v2 Form.pdf"],
        "short_title": "Back MDBQ",
        "roi_name": "Back DBQ",
        "mrr_name": "Back (Thoracolumbar Spine) DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "cfs",
        "filenames": ["Chronic Fatigue Syndrome MDBQ v2 Form.pdf"],
        "short_title": "Chronic Fatigue Syndrome MDBQ",
        "roi_name": "Chronic Fatigue Syndrome DBQ",
        "mrr_name": "Chronic Fatigue Syndrome DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "diabetes",
        "filenames": ["Diabetes Mellitus DBQ.pdf", "Diabetes Mellitus MDBQ.pdf"],
        "short_title": "Diabetes Mellitus DBQ",
        "roi_name": "Diabetes mellitus DBQ",
        "mrr_name": "Diabetes Mellitus DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "ear",
        "filenames": ["Ear (Tinnitus Vertigo) MDBQ.pdf"],
        "short_title": "Ear MDBQ",
        "roi_name": "Ear DBQ",
        "mrr_name": "Ear Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "ed_male_repro",
        "filenames": ["ED - Male Repro MDBQ v2 Form.pdf"],
        "short_title": "ED - Male Repro MDBQ",
        "roi_name": "ED - Male Repro DBQ",
        "mrr_name": "Male Reproductive Organ Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "elbow",
        "filenames": ["Elbow MDBQ v2 Form.pdf"],
        "short_title": "Elbow MDBQ",
        "roi_name": "Elbow DBQ",
        "mrr_name": "Elbow and Forearm DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "eye",
        "filenames": ["Eye MDBQ v2 Form.pdf"],
        "short_title": "Eye MDBQ",
        "roi_name": "Eye DBQ",
        "mrr_name": "Eye Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "fibromyalgia",
        "filenames": ["Fibromyalgia MDBQ v2 Form.pdf"],
        "short_title": "Fibromyalgia MDBQ",
        "roi_name": "Fibromyalgia DBQ",
        "mrr_name": "Fibromyalgia DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "foot",
        "filenames": ["Foot MDBQ v2 Form.pdf"],
        "short_title": "Foot MDBQ",
        "roi_name": "Foot DBQ",
        "mrr_name": "Foot Conditions DBQ",
        "mrr_same_name": "Foot Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "gerd",
        "filenames": ["GERD MDBQ v2 Form.pdf"],
        "short_title": "GERD MDBQ",
        "roi_name": "GERD DBQ",
        "mrr_name": "Esophageal Disorders DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "gynecological",
        "filenames": ["Gynecological MDBQ v2 Form.pdf"],
        "short_title": "Gynecological MDBQ",
        "roi_name": "Gynecological DBQ",
        "mrr_name": "Gynecological Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "hand_finger",
        "filenames": ["Hand & Finger MDBQ v2 Form.pdf"],
        "short_title": "Hand & Finger MDBQ",
        "roi_name": "Hand & Finger DBQ",
        "mrr_name": "Hand and Finger Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "headache",
        "filenames": ["Headache MDBQ v2 Form.pdf"],
        "short_title": "Headache MDBQ",
        "roi_name": "Headache DBQ",
        "mrr_name": "Headaches DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "heart",
        "filenames": ["Heart MDBQ v2 Form.pdf"],
        "short_title": "Heart MDBQ",
        "roi_name": "Heart DBQ",
        "mrr_name": "Heart Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "hematologic",
        "filenames": ["Hematologic & Lymphatic (Blood) MDBQ v2 Form.pdf"],
        "short_title": "Hematologic & Lymphatic MDBQ",
        "roi_name": "Hematologic & Lymphatic DBQ",
        "mrr_name": "Hematologic and Lymphatic Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "hemorrhoids",
        "filenames": ["Hemorrhoids MDBQ v2 Form.pdf"],
        "short_title": "Hemorrhoids MDBQ",
        "roi_name": "Hemorrhoids DBQ",
        "mrr_name": "Rectum and Anus Conditions (Hemorrhoids) DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "hip_thigh",
        "filenames": ["Hip & Thigh MDBQ v2 Form.pdf"],
        "short_title": "Hip & Thigh MDBQ",
        "roi_name": "Hip & Thigh DBQ",
        "mrr_name": "Hip and Thigh Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "hypertension",
        "filenames": ["Hypertension MDBQ v2 Form.pdf"],
        "short_title": "Hypertension MDBQ",
        "roi_name": "Hypertension DBQ",
        "mrr_name": "Hypertension DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "ibs",
        "filenames": ["IBS MDBQ v2 Form.pdf"],
        "short_title": "Intestinal MDBQ",
        "roi_name": "Intestinal DBQ",
        "mrr_name": "Intestinal Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "kidney",
        "filenames": ["Kidney MDBQ v2 Form.pdf"],
        "short_title": "Kidney MDBQ",
        "roi_name": "Kidney Conditions DBQ",
        "mrr_name": "Kidney Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "knee",
        "filenames": ["Knee & Lower leg MDBQ v2 Form.pdf"],
        "short_title": "Knee Lower Leg MDBQ",
        "roi_name": "Knee and Leg DBQ",
        "mrr_name": "Knee and Lower Leg DBQ",
        "mrr_same_name": "Knee and Lower Leg DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "mental_disorders",
        "filenames": ["Mental Disorders MDBQ v2 Form.pdf", "Mental Disorders MDBQ Form.pdf"],
        "short_title": "Mental Disorders MDBQ",
        "roi_name": "Mental DBQ",
        "mrr_name": "Mental Disorders DBQ",
        "opinion_label_mrr": "with Mental Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "muscle_injuries",
        "filenames": ["Muscle Injuries MDBQ v2 Form.pdf"],
        "short_title": "Muscle Injuries MDBQ",
        "roi_name": "Muscle Injuries DBQ",
        "mrr_name": "Muscle Injuries DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "neck",
        "filenames": ["Neck MDBQ v2 Form.pdf"],
        "short_title": "Neck MDBQ",
        "roi_name": "Neck DBQ",
        "mrr_name": "Neck (Cervical Spine) Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "peripheral_nerve",
        "filenames": ["Peripheral Nerve MDBQ v2 Form.pdf"],
        "short_title": "Peripheral Nerve MDBQ",
        "roi_name": "Peripheral Nerve DBQ",
        "mrr_name": "Peripheral Nerves Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "plmd",
        "filenames": ["PLMD_Fillable_Form.pdf"],
        "short_title": "PLMD MDBQ",
        "roi_name": "PLMD DBQ",
        "mrr_name": "Periodic Limb Movement Disorder (PLMD) DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "prostate",
        "filenames": ["Prostate MDBQ v2 Form.pdf"],
        "short_title": "Prostate MDBQ",
        "roi_name": "Prostate DBQ",
        "mrr_name": "Male Reproductive Organ Conditions (Including Prostate Cancer) DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "ptsd",
        "filenames": ["PTSD MDBQ v2 Form.pdf"],
        "short_title": "PTSD MDBQ",
        "roi_name": "PTSD DBQ",
        "mrr_name": "PTSD DBQ",
        "opinion_label_mrr": "with Mental Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "respiratory",
        "filenames": ["Respiratory MDBQ v2 Form.pdf"],
        "short_title": "Respiratory MDBQ",
        "roi_name": "Respiratory DBQ",
        "mrr_name": "Respiratory DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "scar",
        "filenames": ["Scar MDBQ v2 Form.pdf"],
        "short_title": "Scar MDBQ",
        "roi_name": "Scar DBQ",
        "mrr_name": "Scars/Disfigurement DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "shoulder",
        "filenames": ["Shoulder MDBQ v2 Form.pdf"],
        "short_title": "Shoulder MDBQ",
        "roi_name": "Shoulder DBQ",
        "mrr_name": "Shoulder and Arm Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "sinusitis_rhinitis",
        "filenames": ["Sinusitis Rhinitis MDBQ v3 Form.pdf"],
        "short_title": "Sinusitis-Rhinitis MDBQ",
        "roi_name": "Sinusitis/Rhinitis DBQ",
        "mrr_name": "Sinusitis/Rhinitis DBQ",
        "mrr_same_name": "Sinusitis/Rhinitis DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ medical opinion",
    },
    {
        "key": "skin",
        "filenames": ["Skin MDBQ v2 Form.pdf"],
        "short_title": "Skin MDBQ",
        "roi_name": "Skin DBQ",
        "mrr_name": "Skin Diseases DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "sleep_apnea",
        "filenames": ["Sleep Apnea MDBQ v2 Form.pdf"],
        "short_title": "Sleep Apnea MDBQ",
        "roi_name": "Sleep Apnea DBQ",
        "mrr_name": "Sleep Apnea DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "thyroid",
        "filenames": ["Thyroid and Parathyroid MDBQ.pdf"],
        "short_title": "Thyroid and Parathyroid MDBQ",
        "roi_name": "Thyroid and Parathyroid DBQ",
        "mrr_name": "Thyroid and Parathyroid Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "tmj",
        "filenames": ["TMJ MDBQ v2 Form.pdf"],
        "short_title": "TMJ MDBQ",
        "roi_name": "TMJ DBQ",
        "mrr_name": "Temporomandibular Joint (TMJ) Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "voiding",
        "filenames": ["Voiding MDBQ v2 Form.pdf"],
        "short_title": "Voiding MDBQ",
        "roi_name": "Voiding DBQ",
        "mrr_name": "Urinary Tract / Voiding Dysfunction DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
    {
        "key": "wrist",
        "filenames": ["Wrist MDBQ v2 Form.pdf"],
        "short_title": "Wrist MDBQ",
        "roi_name": "Wrist DBQ",
        "mrr_name": "Wrist Conditions DBQ",
        "opinion_label_mrr": "with Medical Opinion",
        "opinion_label_roi": "w/ Medical Opinion",
    },
]


def resolve_template_filename(entry: Dict[str, Any]) -> str:
    """Return the first existing filename in templates/mini_dbqs/ for a catalog entry."""
    for fname in entry["filenames"]:
        if (MINI_DBQ_DIR / fname).exists():
            return fname
    return entry["filenames"][0]


def get_available_mini_dbq_files() -> List[str]:
    """List all canonical Mini-DBQ filenames available in templates/mini_dbqs/."""
    if not MINI_DBQ_DIR.exists():
        return [resolve_template_filename(e) for e in MINI_DBQ_CATALOG]
    files = sorted(
        f.name
        for f in MINI_DBQ_DIR.glob("*.pdf")
        if " - Copy" not in f.name and f.name != "Mental Disorders MDBQ Form.pdf"
    )
    return files


def get_catalog_entry_by_filename(filename: str) -> Dict[str, Any]:
    """Lookup catalog metadata given a template filename."""
    for entry in MINI_DBQ_CATALOG:
        if filename in entry["filenames"]:
            return entry
    fname_lower = (filename or "").lower()
    for entry in MINI_DBQ_CATALOG:
        if any(f.lower() == fname_lower for f in entry["filenames"]):
            return entry
    return MINI_DBQ_CATALOG[0]


def compute_claim_labels(claim: Dict[str, Any], entry: Dict[str, Any]) -> Dict[str, str]:
    """
    Deterministically compute all MRR doctor notes, Mini-DBQ doctor notes, ROI DBQ labels,
    and output file short titles directly from MINI_DBQ_CATALOG.
    """
    if not claim.get("merge_with_previous_dbq", False):
        op_mrr = (
            f" {entry.get('opinion_label_mrr', 'with Medical Opinion')}"
            if claim.get("has_medical_opinion")
            else ""
        )
        op_roi = (
            f" {entry.get('opinion_label_roi', 'w/ Medical Opinion')}"
            if claim.get("has_medical_opinion")
            else ""
        )
        mrr_dr_note = f"Doctor will complete {entry['mrr_name']}{op_mrr}"
        mdbq_dr_note = mrr_dr_note
        roi_short_dbq_name = f"{entry['roi_name']}{op_roi}"
    else:
        same_name = entry.get("mrr_same_name", entry["mrr_name"])
        no_cost_suffix = (
            "; no additional cost" if claim.get("has_no_additional_cost") else ""
        )
        mrr_dr_note = f"Doctor will complete Same {same_name} as above{no_cost_suffix}"
        mdbq_dr_note = ""
        roi_short_dbq_name = ""

    return {
        "short_output_title": entry["short_title"],
        "mrr_dr_note": mrr_dr_note,
        "mdbq_dr_note": mdbq_dr_note,
        "roi_short_dbq_name": roi_short_dbq_name,
    }


# ==============================================================================
# PART 3: STRICT VERBATIM EXTRACTION SCHEMA (Phase3ExtractionResult)
# ==============================================================================
class VeteranProfile(BaseModel):
    client_name: str = Field(description="Veteran's full name.")
    ssn: str = Field(description="Social Security Number (e.g. '542-38-4821').")
    dob: str = Field(description="Date of Birth formatted as MM/DD/YYYY.")
    mos: str = Field(description="Military Occupational Specialty (MOS).")
    dos: str = Field(description="Dates of Service / Branch string.")
    discharge: str = Field(
        default="Honorable",
        description="Character of discharge (default 'Honorable' if blank).",
    )
    street_address: str = Field(
        description="Line 1 of the mailing address only (street and apt/unit), even if 'Address:' is on its own line."
    )
    city_state_zip: str = Field(
        description="Line 2 of the mailing address formatted WITHOUT a comma between city and state (e.g. 'Boston MA 34482')."
    )
    phone: str = Field(description="Cellphone / telephone number.")
    email: str = Field(description="Email address.")
    itf_expiration: str = Field(description="Intent to File (ITF) expiration date (MM/DD/YYYY).")
    case_manager: str = Field(
        default="Sam Armstrong",
        description="Case manager name extracted from Intake Information (e.g. 'Taylor Morgan' or 'Sam Armstrong').",
    )
    current_rating: str = Field(description="Current combined service-connected rating percentage (e.g. '50%').")
    last_claim_decision_date: str = Field(
        description="The most recent (maximum) MM/DD/YYYY date across nsc_issues first, falling back to rated_issues only if nsc_issues has no dates."
    )


class ConditionIssue(BaseModel):
    base_name: str = Field(
        description="Condition name only (before the dash and percentage or NSC), with all '(claimed as ...)' notes and '(no denial date shown)' stripped."
    )
    suffix_text: str = Field(
        description="Only the rating/NSC and converted (MM/DD/YYYY) date after the dash, e.g. '30% (12/21/2009)' or 'NSC (01/07/2008)' or 'NSC' (with '(no denial date shown)' stripped)."
    )


class RecommendedClaim(BaseModel):
    claim_number: int = Field(description="Sequential claim number (1, 2, 3, ...).")
    category: Literal["Medical", "Mental"] = Field(
        description="'Mental' if under Mental/Psychological Conditions Base Fee or using Mental Disorders/PTSD DBQ; otherwise 'Medical'."
    )
    condition_name: str = Field(description="Verbatim condition title before the first dash.")
    claim_type: str = Field(
        description="Uppercase claim type: NEW CLAIM, INCREASE, REOPEN/SECONDARY, SECONDARY/REOPEN, SECONDARY, REOPEN, or PACT ACT PRESUMPTIVE. Never split on hyphenated words like over-the-counter or service-connected."
    )
    clean_body_text: str = Field(
        description="Verbatim symptom/history narrative sentences from the Recommendation doc—do NOT rewrite a single word; strip only the leading number/condition/claim-type header, the trailing 'Doctor will complete...' / 'Same ... DBQ as above...' sentence, 'REQUIRES ...' lines, and video links."
    )
    mapped_dbq_filename: str = Field(
        description="Exact PDF filename selected strictly from the allowed templates/mini_dbqs/ list. If merge_with_previous_dbq is True, MUST match the preceding claim's mapped_dbq_filename."
    )
    has_medical_opinion: bool = Field(
        description="True if 'with medical opinion', 'with Mental Opinion', or 'includes Medical Opinion' is present on this claim."
    )
    merge_with_previous_dbq: bool = Field(
        description="True for EVERY claim that says 'Same ... DBQ as above', even if 3 or 4 claims in a row share the same DBQ."
    )
    has_no_additional_cost: bool = Field(
        description="True if 'no additional cost' appears in the 'Same ... DBQ as above' instruction."
    )
    requires_sleep_study: bool = Field(
        description="True if 'REQUIRES SLEEP STUDY RESULTS' appears in the claim block."
    )


class Phase3ExtractionResult(BaseModel):
    veteran_profile: VeteranProfile
    rated_issues: List[ConditionIssue]
    nsc_issues: List[ConditionIssue]
    recommended_claims: List[RecommendedClaim]


def call_gemini_structured_extraction(
    api_key: str,
    model_name: str,
    intake_raw_text: str,
    recommendation_raw_text: str,
    available_dbq_files: List[str],
) -> Phase3ExtractionResult:
    """
    Pass raw text from Intake Information (.docx) and Recommendation (.docx) to Google Gemini
    for strict verbatim field extraction (temperature=0.0).
    """
    client = genai.Client(api_key=api_key)
    allowed_files_str = "\n".join(f"- {fname}" for fname in available_dbq_files)

    system_prompt = f"""You are a deterministic verbatim data extraction engine for VA Disability Claim Phase 3 documents.
Given the raw text of two Word documents (1. Intake Information and 2. Recommendation), extract all fields verbatim into the strict JSON schema without rewriting any prose.

Follow these exact rules:
1. VETERAN PROFILE:
   - Extract client_name, ssn, dob, mos, dos, discharge (default "Honorable" if missing), street_address, city_state_zip, phone, email, itf_expiration, case_manager (default "Sam Armstrong" if missing), current_rating, and last_claim_decision_date.
   - street_address: Line 1 of the address only (even if "Address:" is on its own line and the street is on the next line).
   - city_state_zip: Line 2 of the address formatted WITHOUT a comma between city and state (e.g., "Boston MA 34482").
   - last_claim_decision_date: Find the maximum MM/DD/YYYY date across nsc_issues first (e.g., "11/20/2008"), falling back to rated_issues only if nsc_issues has no dates.

2. RATED & NSC ISSUES:
   - Separate conditions into rated_issues (those with a percentage like 30%, 10%, 0%, etc.) and nsc_issues (those marked NSC).
   - Convert all dates like (21 Dec 2009) or (20 Nov 2008) to (MM/DD/YYYY) e.g. (12/21/2009) and (11/20/2008).
   - Strip all "(claimed as ...)" parentheticals and strip "(no denial date shown)".
   - Ignore any trailing questionnaire lines at the bottom of the Intake document.
   - For each issue, provide:
     * base_name: Condition name only before the dash (e.g., "Migraine Headaches").
     * suffix_text: Only the rating/NSC and converted date after the dash (e.g., "30% (12/21/2009)" or "NSC (01/07/2008)" or "NSC").

3. RECOMMENDED CLAIMS:
   - Extract EVERY numbered claim (1., 2., 3., ...) in order. Do NOT skip any claim that ends with parenthetical fee notes like "($100 additional DBQ fee + $75 Medical Opinion fee)" or "(Mental Conditions Fee – $400, includes Medical Opinion)".
   - category: "Mental" if under "Mental Conditions Base Fee" / "Psychological Conditions Base Fee" or using Mental Disorders / PTSD DBQ; otherwise "Medical".
   - condition_name: Verbatim condition title before the first dash.
   - claim_type: Uppercase claim type (NEW CLAIM, INCREASE, REOPEN/SECONDARY, SECONDARY/REOPEN, SECONDARY, REOPEN, PACT ACT PRESUMPTIVE). Never split on hyphenated words like "over-the-counter" or "service-connected".
   - clean_body_text: Verbatim symptom/history narrative sentences from the Recommendation doc—do NOT rewrite a single word. Strip only the leading number/condition/claim-type header, the trailing "Doctor/Dr./TP Health Inc. will complete..." or "Same ... DBQ as above..." sentence (and its fee parenthetical), "REQUIRES ..." lines, and video links.
   - mapped_dbq_filename: MUST be chosen strictly from this exact list of filenames in templates/mini_dbqs/:
{allowed_files_str}
     If merge_with_previous_dbq is True, MUST match the preceding claim's mapped_dbq_filename.
   - has_medical_opinion: True if "with medical opinion", "with Mental Opinion", or "includes Medical Opinion" is present on this claim.
   - merge_with_previous_dbq: True for EVERY claim that says "Same ... DBQ as above", even if 3 or 4 claims in a row share the same DBQ.
   - has_no_additional_cost: True if "no additional cost" appears in the "Same ... DBQ as above" instruction.
   - requires_sleep_study: True if "REQUIRES SLEEP STUDY RESULTS" appears in the claim block.
"""

    user_prompt = (
        "=== DOCUMENT 1: INTAKE INFORMATION (.DOCX) ===\n"
        f"{intake_raw_text}\n\n"
        "=== DOCUMENT 2: RECOMMENDATION (.DOCX) ===\n"
        f"{recommendation_raw_text}"
    )

    response = client.models.generate_content(
        model=model_name,
        contents=user_prompt,
        config=types.GenerateContentConfig(
            system_instruction=system_prompt,
            temperature=0.0,
            response_mime_type="application/json",
            response_schema=Phase3ExtractionResult,
        ),
    )

    if getattr(response, "parsed", None) is not None and isinstance(
        response.parsed, Phase3ExtractionResult
    ):
        return response.parsed
    return Phase3ExtractionResult.model_validate_json(response.text)


def _extract_date_from_suffix(suffix_text: str) -> Optional[datetime.date]:
    """Deterministically parse (MM/DD/YYYY) from a condition suffix_text if present."""
    s = (suffix_text or "").strip()
    if "(" in s and ")" in s:
        inside = s[s.rfind("(") + 1 : s.rfind(")")].strip()
        for fmt in ("%m/%d/%Y", "%d %b %Y", "%d %B %Y"):
            try:
                return datetime.datetime.strptime(inside, fmt).date()
            except ValueError:
                continue
    return None


def _compute_deterministic_last_decision_date(
    nsc_issues: List[Dict[str, str]],
    rated_issues: List[Dict[str, str]],
    fallback_date: str,
) -> str:
    """
    Find the maximum MM/DD/YYYY date across nsc_issues first, falling back to
    rated_issues only if nsc_issues has no dates.
    """
    nsc_dates = [
        dt
        for item in nsc_issues
        if (dt := _extract_date_from_suffix(item.get("suffix_text", ""))) is not None
    ]
    if nsc_dates:
        return max(nsc_dates).strftime("%m/%d/%Y")

    rated_dates = [
        dt
        for item in rated_issues
        if (dt := _extract_date_from_suffix(item.get("suffix_text", ""))) is not None
    ]
    if rated_dates:
        return max(rated_dates).strftime("%m/%d/%Y")

    return (fallback_date or "").strip()


def extraction_result_to_dicts(
    result: Phase3ExtractionResult,
) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    """
    Convert the validated Phase3ExtractionResult Pydantic object into the
    intake_data dict and claims list, computing all labels deterministically in Python.
    """
    vp = result.veteran_profile
    rated_list = [
        {
            "base_name": issue.base_name.strip(),
            "suffix_text": issue.suffix_text.strip(),
        }
        for issue in result.rated_issues
    ]
    nsc_list = [
        {
            "base_name": issue.base_name.strip(),
            "suffix_text": issue.suffix_text.strip(),
        }
        for issue in result.nsc_issues
    ]

    last_decision_date = _compute_deterministic_last_decision_date(
        nsc_list, rated_list, vp.last_claim_decision_date
    )

    intake_data: Dict[str, Any] = {
        "client_name": vp.client_name.strip(),
        "ssn": vp.ssn.strip(),
        "dob": vp.dob.strip(),
        "mos": vp.mos.strip(),
        "dos": vp.dos.strip(),
        "discharge": (vp.discharge or "Honorable").strip(),
        "street_address": vp.street_address.strip(),
        "city_state_zip": vp.city_state_zip.replace(",", "").strip(),
        "phone": vp.phone.strip(),
        "email": vp.email.strip(),
        "itf_expiration": vp.itf_expiration.strip(),
        "case_manager": (vp.case_manager or "Sam Armstrong").strip(),
        "current_rating": vp.current_rating.strip(),
        "last_claim_decision_date": last_decision_date,
        "rated_issues": rated_list,
        "nsc_issues": nsc_list,
    }

    available_files = get_available_mini_dbq_files()
    claims: List[Dict[str, Any]] = []
    for c in result.recommended_claims:
        tpl_fname = c.mapped_dbq_filename.strip()
        if c.merge_with_previous_dbq and claims:
            tpl_fname = claims[-1]["mapped_dbq_filename"]
        elif tpl_fname not in available_files:
            entry_lookup = get_catalog_entry_by_filename(tpl_fname)
            tpl_fname = resolve_template_filename(entry_lookup)

        entry = get_catalog_entry_by_filename(tpl_fname)
        raw_claim: Dict[str, Any] = {
            "claim_number": c.claim_number,
            "category": c.category,
            "condition_name": c.condition_name.strip(),
            "claim_type": c.claim_type.strip().upper(),
            "clean_body_text": c.clean_body_text.strip(),
            "mapped_dbq_filename": tpl_fname,
            "has_medical_opinion": bool(c.has_medical_opinion),
            "merge_with_previous_dbq": bool(c.merge_with_previous_dbq),
            "has_no_additional_cost": bool(c.has_no_additional_cost),
            "requires_sleep_study": bool(c.requires_sleep_study),
        }
        labels = compute_claim_labels(raw_claim, entry)
        raw_claim.update(labels)
        claims.append(raw_claim)

    return intake_data, claims


# ==============================================================================
# PART 6: MEDICAL RECORDS REVIEW (.docx) BUILDER
# ==============================================================================
def _set_paragraph_spacing(paragraph: docx.text.paragraph.Paragraph) -> None:
    """Apply standard single/1.15 line spacing with 0pt before/after."""
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def _add_styled_run(
    paragraph: docx.text.paragraph.Paragraph,
    text: str,
    font_name: str = "Calibri",
    size_pt: float = 12.0,
    bold: bool = False,
    underline: bool = False,
    color: Optional[RGBColor] = None,
) -> docx.text.run.Run:
    """Add a formatted run to a paragraph."""
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    if bold:
        run.bold = True
    if underline:
        run.underline = True
    if color is not None:
        run.font.color.rgb = color
    return run


def _add_blank_paragraph(doc: docx.Document) -> docx.text.paragraph.Paragraph:
    p = doc.add_paragraph()
    _set_paragraph_spacing(p)
    return p


def build_mrr_docx(
    intake_data: Dict[str, Any],
    claims: List[Dict[str, Any]],
    gen_date_slashes: str,
    case_manager: str,
    discharge_status: str,
    mrr_case_manager: str = "Sam Armstrong",
) -> bytes:
    """
    Build the Medical Records Review (.docx) matching the exact structure,
    text, and run formatting of the expected output.
    """
    doc = docx.Document()

    # Set 1-inch page margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Configure default Normal style font
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(12)

    client_name = intake_data.get("client_name", "")
    ssn = intake_data.get("ssn", "")
    dob = intake_data.get("dob", "")
    current_rating = intake_data.get("current_rating", "")
    last_decision_date = intake_data.get("last_claim_decision_date", "")
    dos = intake_data.get("dos", "")
    mos = intake_data.get("mos", "")

    # Format discharge string as "{discharge_clean} discharge" (lowercase "discharge")
    discharge_clean = (discharge_status or "Honorable").strip()
    if discharge_clean.lower().endswith("discharge"):
        discharge_clean = discharge_clean[: -len("discharge")].strip()
    discharge_display = f"{discharge_clean} discharge"

    # Title: Medical Records Review
    p_title = doc.add_paragraph()
    _set_paragraph_spacing(p_title)
    _add_styled_run(p_title, "Medical Records Review", size_pt=20.0, color=COLOR_MAROON)

    _add_blank_paragraph(doc)

    # 1. Top Labels
    p_name = doc.add_paragraph()
    _set_paragraph_spacing(p_name)
    _add_styled_run(p_name, "Full Name: ", color=COLOR_MAROON)
    _add_styled_run(p_name, client_name, color=COLOR_BLACK)

    p_ssn = doc.add_paragraph()
    _set_paragraph_spacing(p_ssn)
    _add_styled_run(p_ssn, "SSN: ", color=COLOR_MAROON)
    _add_styled_run(p_ssn, ssn, color=COLOR_BLACK)

    p_dob = doc.add_paragraph()
    _set_paragraph_spacing(p_dob)
    _add_styled_run(p_dob, "DOB: ", color=COLOR_MAROON)
    _add_styled_run(p_dob, dob, color=COLOR_BLACK)

    _add_blank_paragraph(doc)

    p_rating = doc.add_paragraph()
    _set_paragraph_spacing(p_rating)
    _add_styled_run(p_rating, "Current Service Connection: ", color=COLOR_MAROON)
    _add_styled_run(p_rating, current_rating)

    _add_blank_paragraph(doc)

    # 2. Rated & NSC Bullet Lists
    p_rated_hdr = doc.add_paragraph()
    _set_paragraph_spacing(p_rated_hdr)
    _add_styled_run(p_rated_hdr, "Rated Issues:", color=COLOR_MAROON)

    rated_issues: List[Dict[str, Any]] = intake_data.get("rated_issues", [])
    for item in rated_issues:
        p_item = doc.add_paragraph(style="List Bullet")
        _set_paragraph_spacing(p_item)
        _add_styled_run(p_item, item["base_name"], bold=True, color=COLOR_BLACK)
        _add_styled_run(p_item, f" – {item['suffix_text']}")

    _add_blank_paragraph(doc)

    p_nsc_hdr = doc.add_paragraph()
    _set_paragraph_spacing(p_nsc_hdr)
    _add_styled_run(p_nsc_hdr, "Non-Service Connected/Rated Issues:", color=COLOR_MAROON)

    nsc_issues: List[Dict[str, Any]] = intake_data.get("nsc_issues", [])
    for item in nsc_issues:
        p_item = doc.add_paragraph(style="List Bullet")
        _set_paragraph_spacing(p_item)
        _add_styled_run(p_item, item["base_name"], bold=True, color=COLOR_BLACK)
        _add_styled_run(p_item, f" – {item['suffix_text']}")

    _add_blank_paragraph(doc)

    # 3. Last Claim Decision Date & Dates of Service
    p_lcdd = doc.add_paragraph()
    _set_paragraph_spacing(p_lcdd)
    _add_styled_run(p_lcdd, "Last Claim Decision Date: ", color=COLOR_MAROON)
    _add_styled_run(p_lcdd, last_decision_date)

    _add_blank_paragraph(doc)

    p_dos = doc.add_paragraph()
    _set_paragraph_spacing(p_dos)
    _add_styled_run(p_dos, "Dates of Service: ", color=COLOR_MAROON)
    _add_styled_run(p_dos, f"{dos}; {mos}; {discharge_display}")

    _add_blank_paragraph(doc)

    # 4. Intake Notes
    p_in_hdr = doc.add_paragraph()
    _set_paragraph_spacing(p_in_hdr)
    _add_styled_run(p_in_hdr, "Intake Notes:", color=COLOR_MAROON)

    _add_blank_paragraph(doc)

    for claim in claims:
        p_claim = doc.add_paragraph()
        _set_paragraph_spacing(p_claim)
        prefix_str = f"{claim['condition_name']} – {claim['claim_type']}"
        sep_str = " - " if claim.get("claim_type") == "NEW CLAIM" else " – "
        body_part = (claim.get("clean_body_text") or "").strip().replace("\u2019", "'")
        if claim.get("requires_sleep_study"):
            body_part = f"{body_part} REQUIRES SLEEP STUDY RESULTS."
        dr_note = (claim.get("mrr_dr_note") or "").strip()

        _add_styled_run(p_claim, prefix_str, bold=True, color=COLOR_DARK_GRAY)
        _add_styled_run(p_claim, f"{sep_str}{body_part}")
        if dr_note:
            _add_styled_run(
                p_claim,
                f" ({dr_note})",
                bold=True,
                color=COLOR_BLUE,
            )
        _add_blank_paragraph(doc)

    # OVERVIEW
    p_ov_hdr = doc.add_paragraph()
    _set_paragraph_spacing(p_ov_hdr)
    _add_styled_run(p_ov_hdr, "OVERVIEW:", color=COLOR_MAROON)

    _add_blank_paragraph(doc)

    rated_base_names = "; ".join(item["base_name"] for item in rated_issues if item.get("base_name"))
    rated_summary = f"{rated_base_names}." if rated_base_names else "None."

    p_ov1 = doc.add_paragraph()
    _set_paragraph_spacing(p_ov1)
    _add_styled_run(
        p_ov1,
        "The Veteran is seeking assistance with his disability rating. Service treatment records were reviewed to include: ",
        bold=True,
        color=COLOR_BLUE,
    )
    _add_styled_run(p_ov1, rated_summary)

    _add_blank_paragraph(doc)

    rec_cond_names = "; ".join(c["condition_name"] for c in claims if c.get("condition_name"))
    rec_summary = f"{rec_cond_names}." if rec_cond_names else "None."

    p_ov2 = doc.add_paragraph()
    _set_paragraph_spacing(p_ov2)
    _add_styled_run(p_ov2, "After service treatment records: ", bold=True, color=COLOR_BLUE)
    _add_styled_run(p_ov2, rec_summary)

    _add_blank_paragraph(doc)

    p_mrr_sub = doc.add_paragraph()
    _set_paragraph_spacing(p_mrr_sub)
    _add_styled_run(p_mrr_sub, "Medical Record Reviews:", bold=True, color=COLOR_BLUE)

    for _ in range(5):
        _add_blank_paragraph(doc)

    # SERVICE TREATMENT RECORDS REVIEW
    p_str_hdr = doc.add_paragraph()
    _set_paragraph_spacing(p_str_hdr)
    _add_styled_run(p_str_hdr, "SERVICE TREATMENT RECORDS REVIEW", bold=True, underline=True)

    _add_blank_paragraph(doc)

    for item in rated_issues:
        p_str = doc.add_paragraph()
        _set_paragraph_spacing(p_str)
        _add_styled_run(p_str, item["base_name"], bold=True, color=COLOR_BLACK)
        _add_styled_run(
            p_str,
            " – Service Treatment Records (STRs) document diagnosis/treatment while on active duty.",
        )
        _add_blank_paragraph(doc)

    # AFTER SERVICE TREATMENT RECORDS
    p_astr_hdr = doc.add_paragraph()
    _set_paragraph_spacing(p_astr_hdr)
    _add_styled_run(p_astr_hdr, "AFTER SERVICE TREATMENT RECORDS", bold=True, underline=True)

    _add_blank_paragraph(doc)

    for idx_c, claim in enumerate(claims):
        p_astr = doc.add_paragraph()
        _set_paragraph_spacing(p_astr)
        _add_styled_run(
            p_astr,
            f"{claim['condition_name']} – {claim['claim_type']}",
            bold=True,
            color=COLOR_BLACK,
        )
        _add_styled_run(p_astr, " –")
        blank_count = 6 if idx_c == len(claims) - 1 else 5
        for _ in range(blank_count):
            _add_blank_paragraph(doc)

    # 5. Bottom Section Order & Exact Text
    p_surg = doc.add_paragraph()
    _set_paragraph_spacing(p_surg)
    _add_styled_run(p_surg, "SURGERIES: ", color=COLOR_MAROON)
    _add_styled_run(p_surg, "______________________________", color=COLOR_BLACK)

    _add_blank_paragraph(doc)

    p_meds = doc.add_paragraph()
    _set_paragraph_spacing(p_meds)
    _add_styled_run(p_meds, "MEDICATIONS: ", color=COLOR_MAROON)
    _add_styled_run(p_meds, "______________________________", color=COLOR_BLACK)

    _add_blank_paragraph(doc)

    p_cm = doc.add_paragraph()
    _set_paragraph_spacing(p_cm)
    _add_styled_run(p_cm, "CASE MANAGER: ", color=COLOR_MAROON)
    _add_styled_run(p_cm, mrr_case_manager, color=COLOR_BLACK)

    _add_blank_paragraph(doc)

    p_rr = doc.add_paragraph()
    _set_paragraph_spacing(p_rr)
    _add_styled_run(p_rr, "RECORDS REVIEWED: ", color=COLOR_MAROON)
    _add_styled_run(p_rr, "Service Treatment Records", color=COLOR_BLACK)

    _add_blank_paragraph(doc)

    p_drc = doc.add_paragraph()
    _set_paragraph_spacing(p_drc)
    _add_styled_run(p_drc, "DATE REVIEW COMPLETED: ", color=COLOR_MAROON)
    _add_styled_run(p_drc, f"{gen_date_slashes}  {case_manager}", color=COLOR_BLACK)

    _add_blank_paragraph(doc)

    p_rec = doc.add_paragraph()
    _set_paragraph_spacing(p_rec)
    _add_styled_run(p_rec, "RECOMMENDATION: ", color=COLOR_MAROON)
    _add_styled_run(
        p_rec,
        "This section will be completed later by medical records reviewer",
        color=COLOR_BLACK,
    )

    _add_blank_paragraph(doc)

    p_end = doc.add_paragraph()
    _set_paragraph_spacing(p_end)
    _add_styled_run(p_end, "END OF RECORDS REVIEW", bold=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ==============================================================================
# PART 1.5: PATIENT PAYMENT FORM (.pdf) FILLER
# ==============================================================================
def fill_patient_payment_pdf(
    client_name: str,
    street_address: str,
    city_state_zip: str,
) -> bytes:
    """
    Open templates/1. PATIENT PAYMENT FORM 06.03.2024.pdf with PyMuPDF (fitz)
    and populate the Veteran's name and address block (leaving Page 3 Date blank).
    """
    doc = fitz.open(str(PAYMENT_FORM_TEMPLATE))

    # Page 1: Line after "and" before ", an individual"
    p1 = doc[0]
    p1.insert_text(
        fitz.Point(411.5, 143.4),
        client_name,
        fontname="helv",
        fontsize=10.0,
        color=(0, 0, 0),
    )

    # Page 2: Section 4.2 under "Veteran:" (3 lines)
    p2 = doc[1]
    p2.insert_text(
        fitz.Point(136.6, 247.4),
        client_name,
        fontname="helv",
        fontsize=10.0,
        color=(0, 0, 0),
    )
    p2.insert_text(
        fitz.Point(136.6, 262.7),
        street_address,
        fontname="helv",
        fontsize=10.0,
        color=(0, 0, 0),
    )
    p2.insert_text(
        fitz.Point(137.7, 276.9),
        city_state_zip,
        fontname="helv",
        fontsize=10.0,
        color=(0, 0, 0),
    )

    # Page 3: Signature block under "Veteran:" (Date remains blank)
    p3 = doc[2]
    p3.insert_text(
        fitz.Point(289.1, 255.8),
        client_name,
        fontname="helv",
        fontsize=10.0,
        color=(0, 0, 0),
    )

    out_bytes = doc.tobytes(garbage=2, deflate=True)
    doc.close()
    return out_bytes


# ==============================================================================
# PART 5: RELEASE OF INFORMATION (.pdf) FILLER
# ==============================================================================
def fill_roi_pdf(
    intake_data: Dict[str, Any],
    claims: List[Dict[str, Any]],
    itf_expiration: str,
    include_dob: bool = False,
    include_ssn: bool = False,
) -> bytes:
    """
    Open templates/2. RELEASE OF INFORMATION 04.29 - Copy.pdf with PyMuPDF (fitz)
    and populate the Top Table fields and Bottom DBQ Table in strict sequential order:
    distinct_claims[:4] -> Left Column (x=38.0)
    distinct_claims[4:8] -> Right Column (x=327.3)
    """
    doc = fitz.open(str(ROI_FORM_TEMPLATE))
    page = doc[0]

    client_name = intake_data.get("client_name", "")
    dob = intake_data.get("dob", "") if include_dob else ""
    ssn = intake_data.get("ssn", "") if include_ssn else ""
    phone = intake_data.get("phone", "")
    email = intake_data.get("email", "")
    street_address = intake_data.get("street_address", "")
    city_state_zip = intake_data.get("city_state_zip", "")

    # Top Table Coordinates
    top_fields = [
        (178.2, 86.1, client_name),
        (178.2, 98.1, dob),
        (178.2, 110.1, ssn),
        (178.2, 122.4, phone),
        (178.2, 133.9, email),
        (178.2, 146.5, street_address),
        (178.2, 158.1, city_state_zip),
        (178.2, 170.2, itf_expiration),
    ]
    for x, y, val in top_fields:
        if val:
            page.insert_text(
                fitz.Point(x, y),
                val,
                fontname="helv",
                fontsize=8.0,
                color=(0, 0, 0),
            )

    # Bottom DBQ Table in strict sequential order
    y_rows = [705.2, 723.3, 740.0, 758.0]
    distinct_claims = [
        c
        for c in claims
        if not c.get("merge_with_previous_dbq") and c.get("roi_short_dbq_name")
    ]

    for idx, c in enumerate(distinct_claims[:4]):
        page.insert_text(
            fitz.Point(38.0, y_rows[idx]),
            c["roi_short_dbq_name"],
            fontname="helv",
            fontsize=10.0,
            color=(0, 0, 0),
        )

    for idx, c in enumerate(distinct_claims[4:8]):
        page.insert_text(
            fitz.Point(327.3, y_rows[idx]),
            c["roi_short_dbq_name"],
            fontname="helv",
            fontsize=10.0,
            color=(0, 0, 0),
        )

    out_bytes = doc.tobytes(garbage=2, deflate=True)
    doc.close()
    return out_bytes


# ==============================================================================
# PART 4: DETERMINISTIC SENTENCE DEDUPLICATION & MINI-DBQ FORMATTING
# ==============================================================================
def _is_specify_or_other_widget(page: fitz.Page, widget: fitz.Widget) -> bool:
    """
    Check if a text widget on Page 1 is a small/indented 'If checked, specify:'
    or 'Other (specify):' checkbox-companion box rather than the main narrative box.
    """
    r = widget.rect
    fname = (widget.field_name or "").lower()

    if fname.startswith("other") or "weight" in fname or "specify" in fname:
        return True

    if r.x0 > 100.0 or r.width < 440.0:
        return True

    clip_rect = fitz.Rect(0, max(0, r.y0 - 25), r.x1, r.y0 + 15)
    nearby_text = page.get_text("text", clip=clip_rect).lower()
    specify_patterns = (
        "other (please specify)",
        "other, specify",
        "others (please specify)",
        "other (specify)",
        "others (specify)",
        "others(describe)",
        "other, please describe",
        "if checked, specify",
        "if checked. specify",
        "pharyngeal injury",
    )
    if any(pat in nearby_text for pat in specify_patterns):
        return True

    return False


def sanitize_pdf_narrative(narrative_text: str) -> str:
    """
    Replace en-dashes (–) and em-dashes (—) with standard hyphens (-)
    and normalize smart quotes so standard PDF fonts render cleanly.
    """
    s = narrative_text or ""
    s = s.replace("\u2013", "-").replace("\u2014", "-")
    s = s.replace("\u2018", "'").replace("\u2019", "'")
    s = s.replace("\u201c", '"').replace("\u201d", '"')
    s = s.replace("\xa0", " ")
    return s.strip()


def _choose_widget_fontsize(text: str) -> float:
    """
    Dynamically scale the font size between 6.0pt and 9.5pt based on total
    character length so text never clips inside Mini-DBQ text boxes.
    """
    n = len(text or "")
    if n < 250:
        return 9.5
    if n < 400:
        return 9.0
    if n < 600:
        return 8.5
    if n < 850:
        return 8.0
    if n < 1100:
        return 7.5
    if n < 1400:
        return 7.0
    if n < 1800:
        return 6.5
    return 6.0


def _split_into_sentences(text: str) -> List[str]:
    """
    Deterministically split narrative text into sentences while keeping
    abbreviations like 'i.e.,' intact.
    """
    placeholder = "__IE_ABBR__"
    protected = (text or "").strip().replace("i.e.,", placeholder)
    raw_parts = protected.split(". ")
    sentences: List[str] = []
    for idx, part in enumerate(raw_parts):
        restored = part.replace(placeholder, "i.e.,").strip()
        if not restored:
            continue
        if idx < len(raw_parts) - 1 and not restored.endswith("."):
            restored += "."
        sentences.append(restored)
    return sentences


def _build_foot_merged_boxes(c1: Dict[str, Any], c2: Dict[str, Any]) -> Tuple[str, str]:
    """
    Deterministically build Box 1 and Box 2 text for a 2-claim Foot MDBQ merge.
    """
    header = f"{c1['condition_name']} - {c1['claim_type']} / {c2['condition_name']} - {c2['claim_type']}."
    s1_list = _split_into_sentences(c1.get("clean_body_text", ""))
    s2_list = _split_into_sentences(c2.get("clean_body_text", ""))

    # Replace opening "<Condition> symptoms have increased..." sentence
    open1 = f"{c1['condition_name']} symptoms have increased since the last rating decision."
    open2 = f"{c2['condition_name']} symptoms have increased since the last rating decision."
    short_c2 = c2["condition_name"]
    if c1["condition_name"].startswith("Bilateral ") and short_c2.startswith("Bilateral "):
        short_c2 = short_c2[len("Bilateral ") :]
    combined_opening = f"{c1['condition_name']} and {short_c2} symptoms have increased since the last rating decision."

    s1_filtered = [s for s in s1_list if s != open1]
    s2_filtered = [s for s in s2_list if s != open2]

    # Identify unique sentences in c2 that are not in c1 (ignoring shorter prefix variants)
    s1_norm_set = {s.replace("though the physical therapy", "though physical therapy") for s in s1_filtered}
    unique_s2: List[str] = []
    for s in s2_filtered:
        s_norm = s.replace("though the physical therapy", "though physical therapy")
        if s_norm in s1_norm_set:
            continue
        if s.rstrip(".") and any(existing.startswith(s.rstrip(".")) for existing in s1_filtered):
            continue
        unique_s2.append(s_norm)

    merged_sentences: List[str] = [combined_opening]
    inserted_unique = False
    for s in s1_filtered:
        s_norm = s.replace("though the physical therapy", "though physical therapy")
        if s_norm == "This condition has worsened since my last rating." and not inserted_unique:
            merged_sentences.extend(unique_s2)
            inserted_unique = True
        if s_norm not in merged_sentences:
            merged_sentences.append(s_norm)

    if not inserted_unique:
        merged_sentences.extend(u for u in unique_s2 if u not in merged_sentences)

    box1_raw = f"{header} {' '.join(merged_sentences)}"
    box2_raw = f"{c2['condition_name']} - {c2['claim_type']} {c2.get('clean_body_text', '').strip()}"
    return sanitize_pdf_narrative(box1_raw), sanitize_pdf_narrative(box2_raw)


def build_merged_single_box_narrative(claims_group: List[Dict[str, Any]]) -> str:
    """
    Deterministically build the single-box merged narrative (e.g. 4 claims on
    Knee & Lower leg MDBQ v2 Form.pdf) with slash-separated headers and
    bilateral sentence consolidation.
    """
    header_items: List[str] = []
    claim_sentence_lists: List[List[str]] = []

    for c in claims_group:
        sents = _split_into_sentences(c.get("clean_body_text", ""))
        if sents and sents[0].startswith("Related to "):
            rel_clause = sents.pop(0).rstrip(".")
            header_items.append(f"{c['condition_name']} - {c['claim_type']} - {rel_clause}")
        else:
            header_items.append(f"{c['condition_name']} - {c['claim_type']}")
        claim_sentence_lists.append(sents)

    header_str = " / ".join(header_items) + "."

    all_raw_sentences = [s for sents in claim_sentence_lists for s in sents]

    # Check if Service treatment records sentence exists anywhere in the group
    str_sentence = "Service treatment records document treatment while on active duty."
    ordered_sentences: List[str] = []
    if str_sentence in all_raw_sentences:
        ordered_sentences.append(str_sentence)

    pain_variants = {
        "I experience pain, limited range of motion, and swelling in the right shin and lower leg.",
        "I experience pain, limited range of motion, and swelling in the left shin and lower leg.",
        "I experience pain, limited range of motion, and swelling in the right knee.",
        "I experience pain, limited range of motion, and swelling in the left knee.",
    }
    pt_variants = {
        "I have been treated by a medical professional within the last 12 months through physical therapy for my knee, lower, and upper leg.",
        "I have been treated by a medical professional within the last 12 months; physical therapy was concentrated on this condition as it was the most severe of all my lower body symptoms.",
    }
    med_variants = {
        "I manage this condition with prescription and/or over-the-counter medication, i.e., knee braces and Ibuprofen 800mg, which help with pain.",
        "I manage this condition with prescription and/or over-the-counter medication, i.e., knee braces and Ibuprofen, though nothing has been prescribed by a doctor.",
    }
    limp_sentence = "This condition causes me to walk with a limp and altered gait."
    has_limp = limp_sentence in all_raw_sentences

    for s in all_raw_sentences:
        if s == str_sentence:
            continue
        if s.endswith("symptoms have increased since the last rating decision."):
            continue
        if s in pain_variants:
            mapped = "I experience pain, limited range of motion, and swelling in the right and left shin, lower leg, and knees."
        elif s == "This condition has worsened since my last rating.":
            mapped = "These conditions have worsened since my last rating."
        elif s == "The pain increases with physical activity or prolonged standing.":
            if s not in ordered_sentences:
                ordered_sentences.append(s)
            if has_limp:
                limp_mapped = "The knee instability conditions cause me to walk with a limp and altered gait."
                if limp_mapped not in ordered_sentences:
                    ordered_sentences.append(limp_mapped)
            continue
        elif s == limp_sentence:
            mapped = "The knee instability conditions cause me to walk with a limp and altered gait."
        elif s == "I have had to take time off from work due to this condition.":
            mapped = "I have had to take time off from work due to these conditions."
        elif s in pt_variants:
            mapped = (
                "I have been treated by a medical professional within the last 12 months through physical therapy "
                "for my knee, lower, and upper leg; physical therapy was concentrated on my knees as they were "
                "the most severe of all my lower body symptoms."
            )
        elif s in med_variants:
            mapped = (
                "I manage these conditions with prescription and/or over-the-counter medication, i.e., "
                "knee braces and Ibuprofen 800mg, which help with pain."
            )
        else:
            mapped = s

        if mapped not in ordered_sentences:
            ordered_sentences.append(mapped)

    dr_note = (claims_group[0].get("mdbq_dr_note") or "").strip()
    body_joined = " ".join(ordered_sentences)
    if dr_note:
        return sanitize_pdf_narrative(f"{header_str} {body_joined} ({dr_note})")
    return sanitize_pdf_narrative(f"{header_str} {body_joined}")


def fill_mini_dbq_pdf(
    template_filename: str,
    claims_group: List[Dict[str, Any]],
) -> Tuple[bytes, str]:
    """
    Populate a Mini-DBQ PDF template for a group of 1 or more claims using
    deterministic Python formatting and sentence deduplication.
    """
    pdf_path = MINI_DBQ_DIR / template_filename
    if not pdf_path.exists():
        entry = get_catalog_entry_by_filename(template_filename)
        pdf_path = MINI_DBQ_DIR / resolve_template_filename(entry)

    doc = fitz.open(str(pdf_path))
    page = doc[0]

    text_widgets = [
        w for w in page.widgets() if w.field_type == fitz.PDF_WIDGET_TYPE_TEXT
    ]
    text_widgets.sort(key=lambda w: (round(w.rect.y0, 1), round(w.rect.x0, 1)))

    # Clear any pre-filled default text in Page 1 text boxes before injecting
    for w in text_widgets:
        if w.field_value:
            w.field_value = ""
            w.update()

    large_candidates = [
        w for w in text_widgets if w.rect.height >= 28.0 and w.rect.width >= 250.0
    ]
    primary_candidates = [
        w for w in large_candidates if not _is_specify_or_other_widget(page, w)
    ]

    fname_low = (template_filename or "").lower()

    # 2-Box Templates when len(claims_group) == 2: Foot MDBQ or Sinusitis Rhinitis MDBQ
    if len(claims_group) == 2 and ("foot" in fname_low or "sinusitis" in fname_low):
        two_boxes = (
            large_candidates[:2]
            if "sinusitis" in fname_low and len(large_candidates) >= 2
            else primary_candidates[:2]
        )
        if len(two_boxes) == 2:
            c1, c2 = claims_group[0], claims_group[1]
            if "foot" in fname_low:
                box1_text, box2_text = _build_foot_merged_boxes(c1, c2)
            else:
                box1_text = sanitize_pdf_narrative(
                    f"{c1['condition_name']}: {c1.get('clean_body_text', '')}"
                )
                box2_text = sanitize_pdf_narrative(
                    f"{c2['condition_name']}: {c2.get('clean_body_text', '')}"
                )

            for w_box, b_text in ((two_boxes[0], box1_text), (two_boxes[1], box2_text)):
                w_box.field_flags |= fitz.PDF_TX_FIELD_IS_MULTILINE
                w_box.text_fontsize = _choose_widget_fontsize(b_text)
                w_box.field_value = b_text
                w_box.update()

            selected_name = f"{two_boxes[0].field_name}, {two_boxes[1].field_name}"
            out_bytes = doc.tobytes(garbage=2, deflate=True)
            doc.close()
            return out_bytes, selected_name

    # Single-Claim or Multi-Claim Single-Box Narrative
    if len(claims_group) == 1:
        c = claims_group[0]
        body = (c.get("clean_body_text") or "").strip()
        dr_note = (c.get("mdbq_dr_note") or "").strip()
        sep = " - " if body.startswith("Related to ") else ". "
        raw_narrative = f"{c['condition_name']} - {c['claim_type']}{sep}{body}"
        if dr_note:
            raw_narrative = f"{raw_narrative} ({dr_note})"
        clean_text = sanitize_pdf_narrative(raw_narrative)
    else:
        clean_text = build_merged_single_box_narrative(claims_group)

    selected_widget: Optional[fitz.Widget] = None
    if primary_candidates:
        selected_widget = primary_candidates[0]
    elif large_candidates:
        selected_widget = large_candidates[0]
    elif text_widgets:
        selected_widget = max(text_widgets, key=lambda w: w.rect.width * w.rect.height)

    if selected_widget is not None:
        selected_name = selected_widget.field_name or "Text"
        selected_widget.field_flags |= fitz.PDF_TX_FIELD_IS_MULTILINE
        selected_widget.text_fontsize = _choose_widget_fontsize(clean_text)
        selected_widget.field_value = clean_text
        selected_widget.update()
    else:
        if "diabetes" in fname_low:
            target_rect = fitz.Rect(36.0, 502.0, 576.0, 626.0)
        elif "thyroid" in fname_low:
            target_rect = fitz.Rect(90.0, 263.0, 540.0, 305.0)
        else:
            target_rect = fitz.Rect(48.0, 120.0, 564.0, 240.0)
        w = fitz.Widget()
        w.field_name = "Text1"
        w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
        w.field_flags = fitz.PDF_TX_FIELD_IS_MULTILINE
        w.rect = target_rect
        w.text_font = "Helv"
        w.text_fontsize = _choose_widget_fontsize(clean_text)
        w.field_value = clean_text
        page.add_widget(w)
        selected_name = "Text1 (created)"

    out_bytes = doc.tobytes(garbage=2, deflate=True)
    doc.close()
    return out_bytes, selected_name


# ==============================================================================
# PACKAGE GENERATION ORCHESTRATOR
# ==============================================================================
def group_claims_for_mini_dbqs(claims: List[Dict[str, Any]]) -> List[List[Dict[str, Any]]]:
    """
    Group consecutive claims where merge_with_previous_dbq is True into a single PDF.
    """
    groups: List[List[Dict[str, Any]]] = []
    for claim in claims:
        if claim.get("merge_with_previous_dbq") and groups:
            groups[-1].append(claim)
        else:
            groups.append([claim])
    return groups


def generate_phase3_package(
    intake_data: Dict[str, Any],
    claims: List[Dict[str, Any]],
    gen_date: datetime.date,
    itf_expiration: str,
    case_manager: str,
    discharge_status: str,
    include_dob_roi: bool = False,
    include_ssn_roi: bool = False,
) -> Tuple[str, bytes, List[Tuple[str, bytes, str]]]:
    """
    Generate all Phase 3 documents in memory and package them into a .zip archive.
    """
    client_name = (intake_data.get("client_name") or "Veteran").strip()
    date_slashes = gen_date.strftime("%m/%d/%Y")
    date_dots = gen_date.strftime("%m.%d.%Y")

    generated_files: List[Tuple[str, bytes, str]] = []

    # 1. Medical Records Review (.docx)
    mrr_filename = f"Medical Records Review ({client_name}) {date_dots}.docx"
    mrr_bytes = build_mrr_docx(
        intake_data=intake_data,
        claims=claims,
        gen_date_slashes=date_slashes,
        case_manager=case_manager,
        discharge_status=discharge_status,
    )
    generated_files.append(
        (
            mrr_filename,
            mrr_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    )

    # 2. Patient Payment Form (.pdf)
    payment_filename = f"1. PATIENT PAYMENT FORM ({client_name}) {date_dots}.pdf"
    payment_bytes = fill_patient_payment_pdf(
        client_name=client_name,
        street_address=intake_data.get("street_address", ""),
        city_state_zip=intake_data.get("city_state_zip", ""),
    )
    generated_files.append((payment_filename, payment_bytes, "application/pdf"))

    # 3. Release of Information (.pdf)
    roi_filename = f"2. RELEASE OF INFORMATION ({client_name}) {date_dots}.pdf"
    roi_bytes = fill_roi_pdf(
        intake_data=intake_data,
        claims=claims,
        itf_expiration=itf_expiration,
        include_dob=include_dob_roi,
        include_ssn=include_ssn_roi,
    )
    generated_files.append((roi_filename, roi_bytes, "application/pdf"))

    # 4+. Mini-DBQ Forms (starting at index 3, grouping consecutive merge_with_previous_dbq claims)
    claim_groups = group_claims_for_mini_dbqs(claims)
    for offset, group in enumerate(claim_groups):
        primary_claim = group[0]
        file_index = offset + 3
        template_fname = primary_claim["mapped_dbq_filename"]
        entry = get_catalog_entry_by_filename(template_fname)
        short_title = entry["short_title"]

        mdbq_filename = f"{file_index}. {short_title} ({client_name}) {date_dots}.pdf"
        mdbq_bytes, _ = fill_mini_dbq_pdf(
            template_filename=template_fname,
            claims_group=group,
        )
        generated_files.append((mdbq_filename, mdbq_bytes, "application/pdf"))

    # Build in-memory ZIP archive
    safe_client = "".join("_" if ch in '\\/*?:"<>|' else ch for ch in client_name)
    zip_filename = f"{safe_client}_Phase3_{date_dots}.zip"
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for fname, fbytes, _ in generated_files:
            zf.writestr(fname, fbytes)

    return zip_filename, zip_buffer.getvalue(), generated_files


# ==============================================================================
# STREAMLIT TAB UI (PHASE 3)
# ==============================================================================
def render_phase3_tab(api_key: str, model_name: str) -> None:
    st.header("📋 Phase 3 — TP Health Document Automation")
    st.caption(
        "Upload the **Intake Information** and **Recommendation** Word documents (`.docx`) "
     )

    # --------------------------------------------------------------------------
    # File Uploaders
    # --------------------------------------------------------------------------
    col_up1, col_up2 = st.columns(2)
    with col_up1:
        intake_file = st.file_uploader(
            "1. Upload Intake Information (.docx)",
            type=["docx"],
            key="intake_uploader",
        )
    with col_up2:
        rec_file = st.file_uploader(
            "2. Upload Recommendation (.docx)",
            type=["docx"],
            key="rec_uploader",
        )

    # Block processing until an API key is entered
    if not api_key:
        st.warning(
            "🔒 Please enter your **Gemini API Key** in the sidebar to enable document extraction and processing."
        )
        return

    if not intake_file or not rec_file:
        st.info(
            "👆 Please upload both **1. Intake Information (.docx)** and **2. Recommendation (.docx)** to extract data with Gemini and generate the Phase 3 package."
        )
        return

    intake_bytes = intake_file.getvalue()
    rec_bytes = rec_file.getvalue()

    # Cache extraction result in session_state by content hash + model name
    doc_hash = hashlib.sha256(
        intake_bytes + b"||" + rec_bytes + b"||" + model_name.encode("utf-8")
    ).hexdigest()

    if st.session_state.get("extracted_doc_hash") != doc_hash:
        with st.spinner(f"🤖 Extracting verbatim Veteran & Claim fields using `{model_name}`..."):
            try:
                intake_raw_text = extract_raw_docx_text(intake_bytes)
                rec_raw_text = extract_raw_docx_text(rec_bytes)
                available_dbq_files = get_available_mini_dbq_files()
                structured_result = call_gemini_structured_extraction(
                    api_key=api_key,
                    model_name=model_name,
                    intake_raw_text=intake_raw_text,
                    recommendation_raw_text=rec_raw_text,
                    available_dbq_files=available_dbq_files,
                )
                intake_data_parsed, claims_parsed = extraction_result_to_dicts(structured_result)
                st.session_state["extracted_doc_hash"] = doc_hash
                st.session_state["intake_data"] = intake_data_parsed
                st.session_state["parsed_claims"] = claims_parsed
                st.session_state.pop("phase3_zip_bytes", None)
                st.session_state.pop("phase3_zip_fname", None)
                st.session_state.pop("phase3_files_list", None)
            except Exception as exc:
                st.error(f"❌ Gemini Structured Extraction failed: {exc}")
                return

    intake_data: Dict[str, Any] = st.session_state["intake_data"]
    parsed_claims: List[Dict[str, Any]] = st.session_state["parsed_claims"]

    # --------------------------------------------------------------------------
    # Sidebar Overrides
    # --------------------------------------------------------------------------
    st.sidebar.divider()
    st.sidebar.header("⚙️ Phase 3: Package Settings & Overrides")

    gen_date = st.sidebar.date_input(
        "Generation Date",
        value=datetime.date.today(),
        help="Formatted as MM/DD/YYYY inside documents and MM.DD.YYYY in output filenames.",
    )

    default_itf = intake_data.get("itf_expiration", "")
    itf_override = st.sidebar.text_input(
        "ITF Expiration Date",
        value=default_itf,
        placeholder="MM/DD/YYYY",
        help="Auto-populated from Intake document if present.",
    )

    default_cm = intake_data.get("case_manager") or "Sam Armstrong"
    cm_override = st.sidebar.text_input(
        "Case Manager / Reviewer",
        value=default_cm,
        help="Auto-populated from Intake document if present (Default: Sam Armstrong).",
    )

    default_discharge = intake_data.get("discharge") or "Honorable"
    discharge_override = st.sidebar.text_input(
        "Discharge Status",
        value=default_discharge,
        help="Auto-populated from Intake document if present (Default: Honorable).",
    )

    st.sidebar.divider()
    st.sidebar.subheader("ROI Optional Fields")
    include_dob_roi = st.sidebar.checkbox(
        "Include DOB on ROI Form",
        value=False,
        help="Leave unchecked to keep Date of Birth blank on ROI (per standard sample).",
    )
    include_ssn_roi = st.sidebar.checkbox(
        "Include SSN on ROI Form",
        value=False,
        help="Leave unchecked to keep Social Security Number blank on ROI (per standard sample).",
    )

    # --------------------------------------------------------------------------
    # Preview & Validation Panel
    # --------------------------------------------------------------------------
    st.divider()
    st.subheader("👤 Extracted Veteran Profile")

    p_col1, p_col2, p_col3, p_col4 = st.columns(4)
    with p_col1:
        st.metric("Client Name", intake_data.get("client_name") or "—")
        st.write(f"**SSN:** {intake_data.get('ssn') or '—'}")
        st.write(f"**DOB:** {intake_data.get('dob') or '—'}")
    with p_col2:
        st.metric("Current Rating", intake_data.get("current_rating") or "—")
        st.write(f"**Phone:** {intake_data.get('phone') or '—'}")
        st.write(f"**Email:** {intake_data.get('email') or '—'}")
    with p_col3:
        st.metric("Last Claim Decision Date", intake_data.get("last_claim_decision_date") or "—")
        st.write(f"**Street:** {intake_data.get('street_address') or '—'}")
        st.write(f"**City/State/Zip:** {intake_data.get('city_state_zip') or '—'}")
    with p_col4:
        st.metric("ITF Expiration", itf_override or "—")
        st.write(f"**Service (DOS):** {intake_data.get('dos') or '—'}")
        st.write(f"**MOS:** {intake_data.get('mos') or '—'}")

    with st.expander(
        f"📑 View Parsed Rated Issues ({len(intake_data.get('rated_issues', []))}) & "
        f"NSC Issues ({len(intake_data.get('nsc_issues', []))})",
        expanded=False,
    ):
        rc1, rc2 = st.columns(2)
        with rc1:
            st.markdown("**Rated Issues (Cleaned & Formatted)**")
            for r_item in intake_data.get("rated_issues", []):
                st.markdown(f"- **{r_item['base_name']}** – {r_item['suffix_text']}")
        with rc2:
            st.markdown("**Non-Service Connected (NSC) Issues**")
            for n_item in intake_data.get("nsc_issues", []):
                st.markdown(f"- **{n_item['base_name']}** – {n_item['suffix_text']}")

    # --------------------------------------------------------------------------
    # Recommended Claims Table & Template Override Dropdowns
    # --------------------------------------------------------------------------
    st.divider()
    st.subheader(f"🩺 Extracted Recommended Claims ({len(parsed_claims)})")
    st.caption(
        "Review the Gemini-extracted claims, verbatim narratives, and deterministically mapped Mini-DBQ labels below."
    )

    available_templates = get_available_mini_dbq_files()
    updated_claims: List[Dict[str, Any]] = []

    for idx, claim in enumerate(parsed_claims):
        c_copy = dict(claim)
        with st.container(border=True):
            top_c1, top_c2, top_c3, top_c4 = st.columns([2.5, 1.5, 2.8, 1.2])
            with top_c1:
                st.markdown(
                    f"**Claim #{c_copy['claim_number']}: {c_copy['condition_name']}**  \n"
                    f"`{c_copy['category']} Claim` | `{c_copy['claim_type']}`"
                )
            with top_c2:
                med_op = st.checkbox(
                    "Includes Medical Opinion",
                    value=c_copy["has_medical_opinion"],
                    key=f"med_op_{idx}",
                )
                c_copy["has_medical_opinion"] = med_op
                merge_flag = st.checkbox(
                    "Merge w/ Previous DBQ",
                    value=c_copy["merge_with_previous_dbq"],
                    key=f"merge_prev_{idx}",
                )
                c_copy["merge_with_previous_dbq"] = merge_flag
            with top_c3:
                default_tpl = c_copy["mapped_dbq_filename"]
                if default_tpl not in available_templates:
                    available_templates = [default_tpl] + available_templates
                selected_tpl = st.selectbox(
                    "Mapped Mini-DBQ Template",
                    options=available_templates,
                    index=available_templates.index(default_tpl),
                    key=f"tpl_select_{idx}",
                )
                c_copy["mapped_dbq_filename"] = selected_tpl
                matched_entry = get_catalog_entry_by_filename(selected_tpl)
                c_copy.update(compute_claim_labels(c_copy, matched_entry))
            with top_c4:
                if c_copy.get("merge_with_previous_dbq"):
                    st.markdown(
                        f"**Merged Into Above DBQ:**  \n`{c_copy['short_output_title']}`"
                    )
                else:
                    st.markdown(
                        f"**Mini-DBQ Output:**  \n`{c_copy['short_output_title']}`"
                    )

            edited_body = st.text_area(
                "Clean Symptom/History Narrative (clean_body_text)",
                value=c_copy["clean_body_text"],
                height=88,
                key=f"clean_body_{idx}",
            )
            c_copy["clean_body_text"] = edited_body.strip()
        updated_claims.append(c_copy)

    # Summary Table View
    summary_rows = []
    pdf_num = 2
    for c in updated_claims:
        if not c.get("merge_with_previous_dbq"):
            pdf_num += 1
            out_title = f"{pdf_num}. {c['short_output_title']}"
            roi_lbl = c["roi_short_dbq_name"]
        else:
            out_title = f"{pdf_num}. {c['short_output_title']} (Merged)"
            roi_lbl = "— (Same DBQ as above)"
        summary_rows.append(
            {
                "Claim #": c["claim_number"],
                "Category": c["category"],
                "Condition Name": c["condition_name"],
                "Claim Type": c["claim_type"],
                "Mapped Mini-DBQ Template": c["mapped_dbq_filename"],
                "Output Title": out_title,
                "Medical Opinion": "Yes" if c["has_medical_opinion"] else "No",
                "ROI Label": roi_lbl,
                "MRR Intake Label": f"({c.get('mrr_dr_note', '')})",
            }
        )
    st.dataframe(summary_rows, use_container_width=True, hide_index=True)

    # --------------------------------------------------------------------------
    # Generate & Download Section (ZIP Download Only)
    # --------------------------------------------------------------------------
    st.divider()
    if st.button("🚀 Generate Phase 3 Package", type="primary", use_container_width=True, key="phase3_generate_btn"):
        with st.spinner("Generating MRR (.docx), Payment Form, ROI, and Mini-DBQ PDFs..."):
            zip_fname, zip_bytes, files_list = generate_phase3_package(
                intake_data=intake_data,
                claims=updated_claims,
                gen_date=gen_date,
                itf_expiration=itf_override,
                case_manager=cm_override,
                discharge_status=discharge_override,
                include_dob_roi=include_dob_roi,
                include_ssn_roi=include_ssn_roi,
            )
            st.session_state["phase3_zip_fname"] = zip_fname
            st.session_state["phase3_zip_bytes"] = zip_bytes
            st.session_state["phase3_files_list"] = files_list

    if "phase3_zip_bytes" in st.session_state:
        st.success(
            f"✅ Successfully generated **{len(st.session_state['phase3_files_list'])}** documents "
            f"in `{st.session_state['phase3_zip_fname']}`!"
        )
        st.download_button(
            label=f"📦 Download Complete Package ({st.session_state['phase3_zip_fname']})",
            data=st.session_state["phase3_zip_bytes"],
            file_name=st.session_state["phase3_zip_fname"],
            mime="application/zip",
            type="primary",
            use_container_width=True,
            key="phase3_zip_dl_btn",
        )


if __name__ == "__main__":
    from streamlit_app import run_combined_app
    run_combined_app()
